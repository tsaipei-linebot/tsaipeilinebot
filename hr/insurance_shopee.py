"""蝦皮門市的加退保檔案（2026-09-25 新增，見 HANDOFF.md「人資代傳各所加退保 ＋ 收單後補件
＋ 新增『蝦皮』部門」）。

「蝦皮」部門只有人資會上傳，每天有兩種檔案，格式跟 7 個所的 11 欄範本完全不同，這裡負責
把它們讀成人資彙總表（「全區域加退保紀錄表」）的資料列：

- **E-learning**（`KIND_ELEARNING`）：每一列都在「課程權限開通日期(投保日期)」那天**當天加退**；
  沒有日期的照樣納入、投保日留空。部門/店家＝主要門市，備註留空。
- **離店與實習通報**（`KIND_NOTICE`），兩個分頁，名稱每年換（2026實習 → 2027實習），用
  「結尾是『實習』」「結尾是『離店異動』」找：
  - 實習：只處理 A 欄「備註」是「漏保」或空白的列，依「實習日期1」**加保**（沒日期的照樣納入、
    投保日留空）；備註寫「漏保」或「實習」。
  - 離店異動：只處理「異動類型」是「離店」的列，依「異動日期」**退保**；備註寫「離店」。

欄位一律用**表頭名稱**找（不寫死第幾欄），檔案多一欄少一欄也不會對錯欄。每天的檔案都是
當天的資料（不是累計檔），所以不用依日期篩選。

回傳的每一筆是 dict：`store`（部門/店家）、`name`、`id_number`、`insured_text`（彙總表
「投保日」欄的文字）、`note`。廠商一律是 `VENDOR_NAME`。
"""
import io

import openpyxl

KIND_ELEARNING = "elearning"
KIND_NOTICE = "notice"
KIND_NAMES = {
    KIND_ELEARNING: "E-learning",
    KIND_NOTICE: "離店與實習通報",
}
VENDOR_NAME = "蝦皮門市"


class ShopeeFileError(ValueError):
    """檔案格式不對（找不到分頁或必要的欄位），訊息是給人資看的白話文。"""


def _roc(value) -> str:
    # 延後 import，避免 hr.insurance_excel ↔ 這支互相 import
    from hr.insurance_excel import _parse_date, _to_roc_date

    parsed = _parse_date(value)
    return _to_roc_date(parsed) if parsed else ""


def _text(value) -> str:
    return "" if value is None else str(value).strip()


def _header_index(header_row, sheet_title: str, required: dict) -> dict:
    """{代號: 表頭名稱} → {代號: 第幾欄}。同名欄位取第一個（離店異動有兩個「人員隸屬門市」）。
    表頭名稱比對會去掉前後空白；缺欄位就丟 ShopeeFileError。"""
    names = [_text(c) for c in header_row]
    found, missing = {}, []
    for key, header in required.items():
        if header in names:
            found[key] = names.index(header)
        else:
            missing.append(header)
    if missing:
        raise ShopeeFileError(f"「{sheet_title}」分頁找不到這幾欄：{'、'.join(missing)}。請確認上傳的檔案種類是否選對。")
    return found


def _rows(ws):
    rows = ws.iter_rows(values_only=True)
    header = next(rows, None) or ()
    return header, [r for r in rows if r and any(c is not None and _text(c) for c in r)]


def _cell(row, idx):
    return row[idx] if idx < len(row) else None


def _load(content: bytes):
    try:
        return openpyxl.load_workbook(io.BytesIO(content), data_only=True, read_only=True)
    except Exception:
        raise ShopeeFileError("這份檔案打不開，請確認是 Excel 檔（.xlsx），而且沒有設定密碼。")


def parse_elearning(content: bytes) -> list:
    wb = _load(content)
    ws = wb["E-learning"] if "E-learning" in wb.sheetnames else wb.worksheets[0]
    header, rows = _rows(ws)
    idx = _header_index(header, ws.title, {
        "store": "主要門市",
        "id_number": "身分證字號",
        "name": "姓名",
        "date": "課程權限開通日期(投保日期)",
    })
    result = []
    for row in rows:
        name = _text(_cell(row, idx["name"]))
        id_number = _text(_cell(row, idx["id_number"]))
        if not name and not id_number:
            continue
        roc = _roc(_cell(row, idx["date"]))
        result.append({
            "store": _text(_cell(row, idx["store"])),
            "name": name,
            "id_number": id_number,
            "insured_text": f"{roc}當天加退" if roc else "",
            "note": "",
        })
    return result


def _find_sheet(wb, suffix: str):
    for ws in wb.worksheets:
        if ws.title.strip().endswith(suffix):
            return ws
    raise ShopeeFileError(f"找不到名稱結尾是「{suffix}」的分頁（例如「2026{suffix}」）。請確認上傳的檔案種類是否選對。")


def parse_notice(content: bytes) -> list:
    wb = _load(content)
    result = []

    ws = _find_sheet(wb, "實習")
    header, rows = _rows(ws)
    idx = _header_index(header, ws.title, {
        "remark": "備註",
        "name": "人員姓名",
        "store": "人員隸屬門市",
        "id_number": "身分證字號",
        "date": "實習日期1",
    })
    for row in rows:
        remark = _text(_cell(row, idx["remark"]))
        if remark not in ("", "漏保"):
            continue
        name = _text(_cell(row, idx["name"]))
        if not name:
            continue
        roc = _roc(_cell(row, idx["date"]))
        result.append({
            "store": _text(_cell(row, idx["store"])),
            "name": name,
            "id_number": _text(_cell(row, idx["id_number"])),
            "insured_text": f"{roc}加保" if roc else "",
            "note": "漏保" if remark == "漏保" else "實習",
        })

    ws = _find_sheet(wb, "離店異動")
    header, rows = _rows(ws)
    idx = _header_index(header, ws.title, {
        "date": "異動日期",
        "id_number": "身分證字號",
        "name": "姓名",
        "store": "人員隸屬門市",
        "type": "異動類型",
    })
    for row in rows:
        if _text(_cell(row, idx["type"])) != "離店":
            continue
        roc = _roc(_cell(row, idx["date"]))
        result.append({
            "store": _text(_cell(row, idx["store"])),
            "name": _text(_cell(row, idx["name"])),
            "id_number": _text(_cell(row, idx["id_number"])),
            "insured_text": f"{roc}退保" if roc else "",
            "note": "離店",
        })
    return result


PARSERS = {KIND_ELEARNING: parse_elearning, KIND_NOTICE: parse_notice}


def parse(kind: str, content: bytes) -> list:
    if kind not in PARSERS:
        raise ShopeeFileError("請選擇檔案種類。")
    return PARSERS[kind](content)


def describe(kind: str, rows: list) -> str:
    """上傳成功後給人資看的摘要。"""
    if kind == KIND_ELEARNING:
        no_date = sum(1 for r in rows if not r["insured_text"])
        text = f"E-learning 共 {len(rows)} 筆（當天加退）"
        return text + (f"，其中 {no_date} 筆沒有投保日期，彙總表會留空" if no_date else "")
    intern = sum(1 for r in rows if r["note"] in ("實習", "漏保"))
    leave = sum(1 for r in rows if r["note"] == "離店")
    no_date = sum(1 for r in rows if not r["insured_text"])
    text = f"實習加保 {intern} 筆、離店退保 {leave} 筆"
    return text + (f"，其中 {no_date} 筆沒有日期，彙總表會留空" if no_date else "")
