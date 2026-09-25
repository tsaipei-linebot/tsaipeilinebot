"""每日加退保彙總的 Excel 讀取／輸出，跟 `services/contract_summary_excel.py`
是同一種做法（只依賴 openpyxl，輸出 .xlsx bytes，不落地寫檔案），差別是這裡
除了「輸出」，還要「讀取」部門上傳的原始檔案（下載彙總表時即時解析，見
`hr/insurance_repository.py` 開頭「上傳只存檔案本身」的說明）。

**部門上傳檔案固定是 11 欄格式**（編號/廠商/班別/姓名/身分證/勞保加保
日期/勞保退保日期/勞保追退日期/健保加保月份/眷屬健保/備註）——2026-09-22
跟使用者確認過所有部門都用同一份範本，不用各自判斷格式。

**輸出對照「全區域加退保紀錄表」格式**，但這份格式有幾欄系統完全沒有
資料來源（投保單位、出生年月日、招募人員；班次/級距因為原本含金額
資訊、來源檔案的「班別」欄位不夠完整，一併留空），2026-09-22 跟使用者
確認過直接留空，人資下載後自己手動補上，不是程式漏做。「部門/店家」
這欄來源檔案完全沒有對應資料，改填「這份資料是哪個部門傳的」（也就是
上傳時的部門名稱）——是目前系統唯一有的、意義最接近的資訊。「投保日」
這欄則是唯一能從來源資料自動組出來的欄位，把「勞保加保日期」「勞保
退保日期」轉成民國格式接起來（例：「115.09.19當天加退」），比照原本
範例檔案的寫法。
"""
import datetime
import io

import openpyxl
from openpyxl import Workbook

from hr.insurance_draft_repository import format_time
from hr.storage import download_file

_SOURCE_HEADER = (
    "編號", "廠商", "班別", "姓名", "身分證", "勞保加保日期", "勞保退保日期",
    "勞保追退日期", "健保加保月份", "眷屬健保", "備註",
)

_SUMMARY_HEADER = (
    "編號", "投保單位", "廠商", "部門/店家", "姓名", "身分證字號", "投保日",
    "出生年月日", "勞退追退日期", "班次/級距", "備註", "招募人員",
)

# 防 Excel 公式注入，跟 services/contract_summary_excel.py 同一份防呆邏輯
# 說明——來源檔案裡的姓名/備註等欄位是同仁自由填寫的文字，原封不動接進
# 輸出檔案前要擋掉會被 Excel 當成公式執行的開頭字元。
_FORMULA_TRIGGER_CHARS = ("=", "+", "-", "@")


def _sanitize_cell(value):
    if isinstance(value, str) and value.startswith(_FORMULA_TRIGGER_CHARS):
        return "'" + value
    return value


def _to_roc_date(value: datetime.date) -> str:
    """西元轉民國格式「115.09.19」，跟原本「全區域加退保紀錄表」範例
    檔案的日期寫法一致。"""
    return f"{value.year - 1911}.{value.month:02d}.{value.day:02d}"


def _parse_date(raw):
    """盡量把來源儲存格解析成 date：openpyxl 讀到 Excel 日期格式的儲存格
    本來就會是 datetime/date，這裡額外相容常見的文字日期格式；解析不出來
    回傳 None（呼叫端會保留原始文字，不會憑空猜一個日期）。"""
    if isinstance(raw, datetime.datetime):
        return raw.date()
    if isinstance(raw, datetime.date):
        return raw
    text = (str(raw) if raw is not None else "").strip()
    if not text:
        return None
    for fmt in ("%Y-%m-%d", "%Y/%m/%d"):
        try:
            return datetime.datetime.strptime(text, fmt).date()
        except ValueError:
            continue
    return None


def _cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        parsed = _parse_date(value)
        return _to_roc_date(parsed) if parsed else str(value).strip()
    return str(value).strip()


def _insured_date_summary(insured_raw, withdrawn_raw) -> str:
    """組出彙總表「投保日」欄位的文字。兩個日期都能解析、又是同一天 →
    「115.09.19當天加退」；只有其中一個 → 「115.09.19加保」或「…退保」；
    兩個都能解析但不同天 → 兩個日期用「／」接起來；都解析不出來的話，
    原封不動把原始文字接起來，不會因為格式猜不出來就整欄空白。"""
    insured = _parse_date(insured_raw)
    withdrawn = _parse_date(withdrawn_raw)
    if insured and withdrawn:
        if insured == withdrawn:
            return f"{_to_roc_date(insured)}當天加退"
        return f"{_to_roc_date(insured)}加保／{_to_roc_date(withdrawn)}退保"
    if insured:
        return f"{_to_roc_date(insured)}加保"
    if withdrawn:
        return f"{_to_roc_date(withdrawn)}退保"
    insured_text = (str(insured_raw) if insured_raw is not None else "").strip()
    withdrawn_text = (str(withdrawn_raw) if withdrawn_raw is not None else "").strip()
    return "／".join(t for t in (insured_text, withdrawn_text) if t)


def parse_department_workbook(content: bytes) -> list:
    """讀取部門上傳的 11 欄 Excel，回傳每一列的資料字典（欄位對照
    `_SOURCE_HEADER`）。跳過表頭那一列，也跳過姓名／身分證都是空白的列
    （來源檔案常見的尾端空白列）。欄數比預期少的話，缺的欄位補 None，
    不會因為某一欄漏填就整份解析失敗。"""
    wb = openpyxl.load_workbook(io.BytesIO(content), data_only=True)
    ws = wb.active
    rows = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row is None:
            continue
        padded = list(row) + [None] * (len(_SOURCE_HEADER) - len(row))
        record = dict(zip(_SOURCE_HEADER, padded))
        if not _cell_text(record.get("姓名")) and not _cell_text(record.get("身分證")):
            continue
        rows.append(record)
    return rows


def build_summary_workbook(uploads: list) -> bytes:
    """把已經篩好日期區間的上傳紀錄清單（`hr.insurance_repository.
    list_all_history()` 的回傳格式）即時讀取、攤平成「全區域加退保紀錄表」
    格式的一份 Excel。某個部門的檔案讀不到（GCS 找不到或格式解析失敗）
    就跳過那個部門，不會讓整份彙總表下載失敗。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "全區域加退保紀錄表"
    ws.append(list(_SUMMARY_HEADER))

    seq = 1
    for upload in uploads:
        blob_path = upload.get("blob_path") or ""
        if not blob_path:
            continue
        content, _ = download_file(blob_path)
        if content is None:
            continue
        department = upload.get("department", "")
        if upload.get("kind"):
            # 蝦皮（2026-09-25）：E-learning／離店與實習通報，格式不同，見 hr/insurance_shopee.py
            from hr import insurance_shopee

            try:
                shopee_rows = insurance_shopee.parse(upload["kind"], content)
            except Exception:
                continue
            for record in shopee_rows:
                ws.append([
                    _sanitize_cell(seq), "", insurance_shopee.VENDOR_NAME,
                    _sanitize_cell(record["store"]), _sanitize_cell(record["name"]),
                    _sanitize_cell(record["id_number"]), _sanitize_cell(record["insured_text"]),
                    "", "", "", _sanitize_cell(record["note"]), "",
                ])
                seq += 1
            continue
        try:
            source_rows = parse_department_workbook(content)
        except Exception:
            continue
        for record in source_rows:
            ws.append([
                _sanitize_cell(seq),
                "",
                _sanitize_cell(_cell_text(record.get("廠商"))),
                _sanitize_cell(department),
                _sanitize_cell(_cell_text(record.get("姓名"))),
                _sanitize_cell(_cell_text(record.get("身分證"))),
                _sanitize_cell(_insured_date_summary(record.get("勞保加保日期"), record.get("勞保退保日期"))),
                "",
                _sanitize_cell(_cell_text(record.get("勞保追退日期"))),
                "",
                _sanitize_cell(_cell_text(record.get("備註"))),
                "",
            ])
            seq += 1

    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


# ==========================================
# 加退保暫存區（2026-09-24 新增，見 hr/insurance_draft_repository.py）
# ==========================================
_SOURCE_DATE_HEADERS = {"勞保加保日期", "勞保退保日期", "勞保追退日期"}


def _source_cell(header, value):
    """寫回部門範本格式時，日期欄位盡量寫成真正的 Excel 日期（人資開檔看到
    的是日期、`_parse_date()` 讀回來也認得），其他欄位當文字並擋公式注入。"""
    if header in _SOURCE_DATE_HEADERS:
        parsed = _parse_date(value)
        if parsed:
            return parsed
    if value is None:
        return ""
    if isinstance(value, (datetime.datetime, datetime.date)):
        return value
    return _sanitize_cell(str(value).strip())


def build_department_workbook(rows: list) -> bytes:
    """把資料列（key 是 `_SOURCE_HEADER` 的表頭）組成跟部門上傳範本一樣的 11 欄
    Excel。暫存區「送出給人資」「下載」都用這個，所以產出的檔案
    `parse_department_workbook()` 讀得回來，人資端的彙總表不用改。「編號」
    一律重新從 1 編。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "加退保"
    ws.append(list(_SOURCE_HEADER))
    for seq, row in enumerate(rows, start=1):
        cells = [seq] + [_source_cell(h, row.get(h)) for h in _SOURCE_HEADER[1:]]
        ws.append(cells)
        for col, header in enumerate(_SOURCE_HEADER, start=1):
            if header in _SOURCE_DATE_HEADERS and isinstance(cells[col - 1], datetime.date):
                ws.cell(row=seq + 1, column=col).number_format = "yyyy-mm-dd"
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


def build_draft_records_workbook(drafts: list, status_names: dict, type_name) -> bytes:
    """「加退保操作紀錄」頁的匯出：範本 11 欄之外，多列部門、類型、狀態、誰建立、
    最後一次操作。只是查紀錄用，不會改變任何一筆的狀態。"""
    wb = Workbook()
    ws = wb.active
    ws.title = "加退保操作紀錄"
    ws.append(["部門", "類型", "狀態", "交給人資的日期"] + list(_SOURCE_HEADER[1:]) + ["建立者", "建立時間", "最後操作"])
    for draft in drafts:
        history = draft.get("history") or []
        last = history[-1] if history else {}
        ws.append(
            [
                _sanitize_cell(draft.get("department", "")),
                type_name(draft),
                status_names.get(draft.get("status"), draft.get("status", "")),
                draft.get("sent_work_date", ""),
            ]
            + [_sanitize_cell(str(draft.get(key) or "")) for key in (
                "vendor", "shift", "name", "id_number", "insured_date", "withdrawn_date",
                "recovery_date", "health_month", "dependents", "note",
            )]
            + [
                _sanitize_cell(draft.get("created_by_name", "")),
                format_time(draft.get("created_at")),
                _sanitize_cell(f"{last.get('by_name', '')} {format_time(last.get('at'))}".strip()),
            ]
        )
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()

