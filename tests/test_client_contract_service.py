import os
import sys
import unittest
from datetime import date
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests import _env  # noqa: F401
from tests import _stub_gcp
_stub_gcp.install()

import docx

from services.client_contract_service import (
    CONTRACT_VERSIONS,
    DEFAULT_CONTRACT_VERSION,
    SEVERANCE_PAYER_OPTIONS,
    can_view_submission,
    default_contract_end_date,
    list_visible_submissions,
    render_contract_docx,
    roc_date_string,
)

_PARTY_A = {
    "name": "測試客戶股份有限公司", "representative": "陳大明",
    "address": "台北市信義區忠孝東路一段1號", "tax_id": "12345678", "phone": "02-1234-5678",
}
_PARTY_B = {
    "name": "瑋政有限公司", "representative": "蔡志祥",
    "address": "新北市板橋區文化路二段90號五樓", "tax_id": "68138452", "phone": "(02) 6637-3899",
}


class RocDateStringTests(unittest.TestCase):
    def test_converts_gregorian_to_roc(self):
        self.assertEqual(roc_date_string(date(2026, 1, 1)), "115年01月01日")
        self.assertEqual(roc_date_string(date(2025, 7, 28)), "114年07月28日")

    def test_pads_single_digit_month_and_day(self):
        self.assertEqual(roc_date_string(date(2026, 3, 5)), "115年03月05日")


class DefaultContractEndDateTests(unittest.TestCase):
    def test_defaults_to_december_31_of_given_year(self):
        self.assertEqual(default_contract_end_date(date(2026, 3, 15)), date(2026, 12, 31))

    def test_uses_today_when_not_given(self):
        result = default_contract_end_date()
        self.assertEqual(result.month, 12)
        self.assertEqual(result.day, 31)


class RenderContractDocxTests(unittest.TestCase):
    """render_contract_docx() 會真的讀 assets/ 底下的 master template 檔案、
    用 docxtpl 套版，但不碰 Firestore/GCS，歸類為可以直接測的部分。"""

    def _render(self, **overrides):
        defaults = dict(
            party_a=_PARTY_A,
            party_b=_PARTY_B,
            sign_date=date(2026, 1, 1),
            contract_start_date=date(2026, 1, 1),
            contract_end_date=date(2026, 12, 31),
            replace_notice_days="3",
            severance_payer="乙方",
            remit_day="10",
            hourly_wage="200",
            management_fee="65",
        )
        defaults.update(overrides)
        content = render_contract_docx(**defaults)
        path = "/tmp/_test_client_contract_render.docx"
        with open(path, "wb") as f:
            f.write(content)
        return docx.Document(path)

    def test_no_leftover_jinja_tags(self):
        doc = self._render()
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertNotIn("{{", full_text)
        self.assertNotIn("{%", full_text)

    def test_party_a_and_b_names_appear(self):
        doc = self._render()
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("測試客戶股份有限公司", full_text)
        self.assertIn("瑋政有限公司", full_text)

    def test_party_identity_table_has_all_fields(self):
        doc = self._render()
        table = doc.tables[0]
        a_col_text = "\n".join(p.text for p in table.rows[1].cells[0].paragraphs)
        b_col_text = "\n".join(p.text for p in table.rows[1].cells[1].paragraphs)
        self.assertIn("陳大明", a_col_text)
        self.assertIn("台北市信義區忠孝東路一段1號", a_col_text)
        self.assertIn("12345678", a_col_text)
        self.assertIn("02-1234-5678", a_col_text)
        self.assertIn("蔡志祥", b_col_text)
        self.assertIn("68138452", b_col_text)

    def test_dates_rendered_in_roc_format_without_duplicate_day_character(self):
        doc = self._render()
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("115年01月01日", full_text)
        self.assertIn("115年12月31日", full_text)
        self.assertNotIn("日日", full_text)

    def test_replace_notice_days_and_severance_payer(self):
        doc = self._render(replace_notice_days="5", severance_payer="甲方")
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("甲方得要求乙方須於5日內撤換此人員", full_text)
        self.assertIn("甲方須支付資遣費用及預告工資", full_text)

    def test_remit_day_substituted(self):
        doc = self._render(remit_day="20")
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("甲方應於次月20日前匯款", full_text)

    def test_rate_table_substituted(self):
        doc = self._render(hourly_wage="180", management_fee="55")
        table = doc.tables[1]
        row = table.rows[1]
        self.assertEqual(row.cells[1].text, "180元/hr")
        self.assertEqual(row.cells[3].text, "55元/hr")


class ContractVersionConfigTests(unittest.TestCase):
    def test_default_version_exists_in_registry(self):
        self.assertIn(DEFAULT_CONTRACT_VERSION, CONTRACT_VERSIONS)

    def test_severance_payer_options_are_the_two_parties(self):
        self.assertEqual(set(SEVERANCE_PAYER_OPTIONS), {"甲方", "乙方"})


class CanViewSubmissionTests(unittest.TestCase):
    """跟 dispatch_contract_service 同一套規則：送出者本人/主管/平台管理員
    才看得到。"""

    def test_submitter_can_view_own_record(self):
        viewer = {"username": "bob", "is_platform_admin": False}
        record = {"submitted_by": "bob"}
        self.assertTrue(can_view_submission(viewer, record))

    def test_platform_admin_can_view_anyone(self):
        viewer = {"username": "boss", "is_platform_admin": True}
        record = {"submitted_by": "alice"}
        self.assertTrue(can_view_submission(viewer, record))

    def test_manager_of_submitter_can_view(self):
        viewer = {"username": "carol", "is_platform_admin": False}
        record = {"submitted_by": "alice"}
        with mock.patch("services.client_contract_service.platform_accounts.get_account",
                         return_value={"username": "alice", "manager_usernames": ["carol"]}):
            self.assertTrue(can_view_submission(viewer, record))

    def test_unrelated_account_cannot_view(self):
        viewer = {"username": "dave", "is_platform_admin": False}
        record = {"submitted_by": "alice"}
        with mock.patch("services.client_contract_service.platform_accounts.get_account",
                         return_value={"username": "alice", "manager_usernames": ["carol"]}):
            self.assertFalse(can_view_submission(viewer, record))


class ListVisibleSubmissionsTests(unittest.TestCase):
    def test_platform_admin_sees_all(self):
        viewer = {"username": "boss", "is_platform_admin": True}
        records = [{"submitted_by": "alice"}, {"submitted_by": "bob"}]
        with mock.patch("services.client_contract_service.list_submissions", return_value=records):
            self.assertEqual(list_visible_submissions(viewer), records)

    def test_non_admin_only_sees_own_and_subordinates(self):
        viewer = {"username": "carol", "is_platform_admin": False}
        records = [
            {"id": "1", "submitted_by": "carol"},
            {"id": "2", "submitted_by": "alice"},
            {"id": "3", "submitted_by": "dave"},
        ]

        def _fake_get_account(username):
            accounts = {
                "alice": {"username": "alice", "manager_usernames": ["carol"]},
                "dave": {"username": "dave", "manager_usernames": ["someone_else"]},
            }
            return accounts.get(username)

        with mock.patch("services.client_contract_service.list_submissions", return_value=records):
            with mock.patch("services.client_contract_service.platform_accounts.get_account",
                             side_effect=_fake_get_account):
                visible = list_visible_submissions(viewer)
        self.assertEqual([r["id"] for r in visible], ["1", "2"])


class ConvertDocxToPdfDelegatesToSharedModuleTests(unittest.TestCase):
    def test_delegates_to_shared_conversion_with_client_contract_log_prefix(self):
        from services.client_contract_service import convert_docx_to_pdf
        with mock.patch("services.client_contract_service._convert_docx_to_pdf",
                         return_value=b"%PDF-FAKE") as mock_convert:
            result = convert_docx_to_pdf(b"fake docx bytes")
        self.assertEqual(result, b"%PDF-FAKE")
        mock_convert.assert_called_once_with(b"fake docx bytes", log_prefix="[合約產生器 PDF 轉檔失敗]")


if __name__ == "__main__":
    unittest.main()
