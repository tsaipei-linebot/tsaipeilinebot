import datetime
import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests import _env  # noqa: F401

from services import tabular_upload


def _xlsx(rows: list) -> bytes:
    """用 openpyxl 產一份 xlsx，rows[0] 當表頭。值直接照傳進來的型別寫入，
    這樣才能驗證日期/數字型儲存格（同仁在 Excel 裡打日期，Excel 存的是
    日期型別，不是文字）。"""
    from openpyxl import Workbook

    workbook = Workbook()
    worksheet = workbook.active
    for row in rows:
        worksheet.append(row)
    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()


class CellToTextTests(unittest.TestCase):
    def test_none_becomes_empty_string(self):
        self.assertEqual(tabular_upload.cell_to_text(None), "")

    def test_integral_float_loses_the_decimal_point(self):
        # Excel 的數字一律是浮點數，不處理的話「需求人數」會變成 "3.0" 而解析失敗
        self.assertEqual(tabular_upload.cell_to_text(3.0), "3")

    def test_real_decimal_is_kept(self):
        # 緯經度不能被當成整數處理掉
        self.assertEqual(tabular_upload.cell_to_text(24.98801), "24.98801")

    def test_midnight_datetime_is_treated_as_date_only(self):
        self.assertEqual(tabular_upload.cell_to_text(datetime.datetime(2024, 1, 31)), "2024-01-31")

    def test_datetime_with_time_keeps_the_time(self):
        self.assertEqual(tabular_upload.cell_to_text(datetime.datetime(2024, 1, 31, 9, 5)), "2024-01-31 09:05")

    def test_date_object(self):
        self.assertEqual(tabular_upload.cell_to_text(datetime.date(2024, 1, 31)), "2024-01-31")

    def test_bool_is_not_rendered_as_one_or_zero(self):
        # Python 的 bool 是 int 的子類別，不特別處理會變成 "1"/"0"
        self.assertEqual(tabular_upload.cell_to_text(True), "是")
        self.assertEqual(tabular_upload.cell_to_text(False), "否")


class LooksLikeXlsxTests(unittest.TestCase):
    def test_xlsx_is_detected_by_content_not_filename(self):
        self.assertTrue(tabular_upload.looks_like_xlsx(_xlsx([["姓名"], ["王小明"]])))

    def test_csv_bytes_are_not_mistaken_for_xlsx(self):
        self.assertFalse(tabular_upload.looks_like_xlsx("姓名\n王小明\n".encode("utf-8")))

    def test_empty_is_not_xlsx(self):
        self.assertFalse(tabular_upload.looks_like_xlsx(b""))


class ReadRowsTests(unittest.TestCase):
    def test_empty_content(self):
        rows, error = tabular_upload.read_rows(b"")
        self.assertEqual(rows, [])
        self.assertIn("空", error)

    def test_csv_utf8_with_bom(self):
        rows, error = tabular_upload.read_rows("姓名,電話\n陳堃喆,0912345678\n".encode("utf-8-sig"))
        self.assertIsNone(error)
        self.assertEqual(rows, [{"姓名": "陳堃喆", "電話": "0912345678"}])

    def test_csv_big5_still_decodes(self):
        """舊的 Big5 檔案要繼續讀得動（只是 Big5 放不下的字在存檔時就已經
        變成 ? 了，救不回來——這正是改用 Excel 的原因）。"""
        rows, error = tabular_upload.read_rows("姓名,電話\n王小明,0912345678\n".encode("cp950"))
        self.assertIsNone(error)
        self.assertEqual(rows[0]["姓名"], "王小明")

    def test_xlsx_rows_are_read_as_strings(self):
        content = _xlsx([["姓名", "人數", "到職日期"], ["陳堃喆", 3, datetime.datetime(2024, 1, 31)]])
        rows, error = tabular_upload.read_rows(content)
        self.assertIsNone(error)
        self.assertEqual(rows, [{"姓名": "陳堃喆", "人數": "3", "到職日期": "2024-01-31"}])

    def test_xlsx_short_row_is_padded_not_rejected(self):
        # Excel 尾端沒填的儲存格常常直接不回傳，缺的欄位要補空字串
        content = _xlsx([["姓名", "電話", "備註"], ["王小明", "0912345678"]])
        rows, error = tabular_upload.read_rows(content)
        self.assertIsNone(error)
        self.assertEqual(rows[0]["備註"], "")

    def test_xlsx_with_only_headers_gives_no_rows_but_no_error(self):
        rows, error = tabular_upload.read_rows(_xlsx([["姓名", "電話"]]))
        self.assertIsNone(error)
        self.assertEqual(rows, [])

    def test_broken_xlsx_returns_plain_message_instead_of_raising(self):
        # 開頭像 xlsx（zip）但內容是壞的
        rows, error = tabular_upload.read_rows(b"PK\x03\x04this is not really a workbook")
        self.assertEqual(rows, [])
        self.assertIn("Excel", error)


class HeaderNamesTests(unittest.TestCase):
    def test_empty_rows_give_empty_set(self):
        self.assertEqual(tabular_upload.header_names([]), set())

    def test_names_are_stripped_and_blanks_dropped(self):
        self.assertEqual(tabular_upload.header_names([{" 姓名 ": "x", "": "y", "電話": "z"}]), {"姓名", "電話"})


class BuildTemplateXlsxTests(unittest.TestCase):
    def _load(self, **kwargs):
        import openpyxl

        content = tabular_upload.build_template_xlsx(
            ["廠商", "姓名", "電話"], ["蝦皮三輪", "王小明", "0912345678"], **kwargs
        )
        return content, openpyxl.load_workbook(io.BytesIO(content)).active

    def test_template_has_exactly_header_and_sample_rows(self):
        """曾經踩過的雷：逐一設定每個儲存格的格式會把那些儲存格實際建出來，
        `append()` 就接在它們後面，範例資料被擠到第 1002 列。"""
        _, worksheet = self._load(text_columns=("電話",))
        self.assertEqual(worksheet.max_row, 2)
        self.assertEqual([c.value for c in worksheet[1]], ["廠商", "姓名", "電話"])
        self.assertEqual([c.value for c in worksheet[2]], ["蝦皮三輪", "王小明", "0912345678"])

    def test_text_column_is_formatted_as_text(self):
        # 不設的話 Excel 會把 0912345678 當數字、開頭的 0 直接不見
        _, worksheet = self._load(text_columns=("電話",))
        self.assertEqual(worksheet.column_dimensions["C"].number_format, "@")
        self.assertEqual(worksheet.cell(row=2, column=3).number_format, "@")

    def test_non_text_columns_are_left_alone(self):
        _, worksheet = self._load(text_columns=("電話",))
        self.assertNotEqual(worksheet.column_dimensions["A"].number_format, "@")

    def test_template_can_be_read_back_by_read_rows(self):
        """範本產出來之後要真的讀得回去——這條連起「下載範本 → 填 → 上傳」
        整個流程，是這次改動最重要的一條。"""
        content, _ = self._load(text_columns=("電話",))
        rows, error = tabular_upload.read_rows(content)
        self.assertIsNone(error)
        self.assertEqual(rows, [{"廠商": "蝦皮三輪", "姓名": "王小明", "電話": "0912345678"}])


if __name__ == "__main__":
    unittest.main()
