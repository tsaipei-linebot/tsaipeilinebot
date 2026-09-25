"""蝦皮門市加退保檔案（2026-09-25 新增，hr/insurance_shopee.py）。

測試資料全部是假的（王小明、A123456789…），真的範例檔含個資，沒有放進 repo。"""
import datetime
import io
import os
import sys
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests import _env  # noqa: F401
from tests import _stub_gcp
_stub_gcp.install()

import openpyxl
from fastapi.testclient import TestClient

import main
import platform_accounts
from hr import insurance_shopee as shopee
from hr import insurance_repository as repo
from hr.insurance_excel import build_summary_workbook
from hr.routes import insurance_routes
from tests._fake_firestore import FakeFirestore

XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
HR = {"username": "hr", "name": "HR", "department": "人資部門", "modules": ["hr"], "is_platform_admin": False, "rank": ""}
TAOYUAN = {"username": "t", "name": "T", "department": "桃園所", "modules": [], "is_platform_admin": False, "rank": ""}
D = datetime.datetime


def _book(sheets: dict) -> bytes:
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for row in rows:
            ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


ELEARNING_HEADER = ["建檔日期", "離店日期", "在職狀態", "主要門市", "身分證字號", "姓名", "教育訓練日", "人員填表日期",
                    "課程權限開通日期(投保日期)", "課程類型"]


def elearning_file():
    return _book({"E-learning": [
        ELEARNING_HEADER,
        [D(2026, 9, 20), None, "建檔中", "板橋長江 - 智取店", "A123456789", "王小明", None, None, D(2026, 9, 21), "智取店"],
        [D(2026, 9, 20), D(2026, 9, 30), "已離職", "新莊中平店", "B223456789", "陳小華", None, None, None, "一般店"],
        [None] * 10,
    ]})


INTERN_HEADER = ["備註", "人員姓名", "人員職稱", "人員隸屬門市", "身分證字號", "派遣公司", "實習日期1", "實習日期2", "漏保"]
LEAVE_HEADER = ["異動日期", "派遣公司", "身分證字號", "姓名", "人員隸屬門市", "異動類型", "異動說明", "離店備註",
                "人員隸屬門市", "轉調前職稱", "人員轉入店名", "轉調後職稱"]


def notice_file(year="2026"):
    return _book({
        f"{year}實習": [
            INTERN_HEADER,
            [None, "王小明", "早班時薪", "板橋長江 - 智取店", "A123456789", "材霈有限公司", D(2026, 9, 21), None, None],
            ["漏保", "陳小華", "晚班時薪", "新莊中平店", "B223456789", "材霈有限公司", "2026/09/19", None, "V"],
            ["取消實習", "林小美", "晚班時薪", "新莊中平店", "C123456789", "材霈有限公司", D(2026, 9, 22), None, None],
            ["增加22", "張三", "晚班時薪", "新莊中平店", "D123456789", "材霈有限公司", D(2026, 9, 22), None, None],
            [None, "李四", "晚班時薪", "新莊中平店", "E123456789", "材霈有限公司", None, None, None],
        ],
        f"{year}離店異動": [
            LEAVE_HEADER,
            [D(2026, 9, 19), "材霈有限公司", "F123456789", "趙五", "士林格致 - 智取店", "離店", "", "", "另一家店", "", "", ""],
            [D(2026, 9, 19), "材霈有限公司", "G123456789", "錢六", "新莊中平店", "留停復職", "", "", "", "", "", ""],
            [D(2026, 9, 20), "材霈有限公司", "H123456789", "孫七", "中壢華美 - 智取店", "單位/身分異動", "", "", "", "", "", ""],
        ],
    })


class ParseTests(unittest.TestCase):
    def test_elearning_is_same_day_add_and_remove(self):
        rows = shopee.parse(shopee.KIND_ELEARNING, elearning_file())
        self.assertEqual(rows, [
            {"store": "板橋長江 - 智取店", "name": "王小明", "id_number": "A123456789", "insured_text": "115.09.21當天加退", "note": ""},
            # 沒有日期照樣納入、投保日留空；已離職、課程類型都不管
            {"store": "新莊中平店", "name": "陳小華", "id_number": "B223456789", "insured_text": "", "note": ""},
        ])

    def test_notice_intern_only_blank_or_missed_and_leave_only_leave(self):
        rows = shopee.parse(shopee.KIND_NOTICE, notice_file())
        self.assertEqual([(r["name"], r["insured_text"], r["note"], r["store"]) for r in rows], [
            ("王小明", "115.09.21加保", "實習", "板橋長江 - 智取店"),
            ("陳小華", "115.09.19加保", "漏保", "新莊中平店"),
            ("李四", "", "實習", "新莊中平店"),
            ("趙五", "115.09.19退保", "離店", "士林格致 - 智取店"),  # 兩個「人員隸屬門市」取第一個
        ])

    def test_sheet_names_follow_the_year(self):
        rows = shopee.parse(shopee.KIND_NOTICE, notice_file(year="2027"))
        self.assertEqual(len(rows), 4)

    def test_wrong_kind_gives_a_plain_error(self):
        with self.assertRaises(shopee.ShopeeFileError) as ctx:
            shopee.parse(shopee.KIND_NOTICE, elearning_file())
        self.assertIn("實習", str(ctx.exception))
        with self.assertRaises(shopee.ShopeeFileError) as ctx:
            shopee.parse(shopee.KIND_ELEARNING, notice_file())
        self.assertIn("課程權限開通日期(投保日期)", str(ctx.exception))

    def test_not_an_excel_file(self):
        with self.assertRaises(shopee.ShopeeFileError):
            shopee.parse(shopee.KIND_ELEARNING, b"not excel")

    def test_describe(self):
        self.assertEqual(shopee.describe(shopee.KIND_NOTICE, shopee.parse(shopee.KIND_NOTICE, notice_file())),
                         "實習加保 3 筆、離店退保 1 筆，其中 1 筆沒有日期，彙總表會留空")


class SummaryWorkbookTests(unittest.TestCase):
    def test_shopee_rows_in_summary(self):
        blobs = {"e": elearning_file(), "n": notice_file()}
        uploads = [
            {"department": "蝦皮", "kind": "elearning", "blob_path": "e"},
            {"department": "蝦皮", "kind": "notice", "blob_path": "n"},
        ]
        with mock.patch("hr.insurance_excel.download_file", side_effect=lambda p: (blobs[p], "x")):
            content = build_summary_workbook(uploads)
        ws = openpyxl.load_workbook(io.BytesIO(content)).active
        rows = [[c.value for c in r] for r in ws.iter_rows(min_row=2)]
        self.assertEqual(len(rows), 6)
        first = rows[0]
        # 編號/投保單位/廠商/部門店家/姓名/身分證/投保日/…/備註
        self.assertEqual((first[0], first[2], first[3], first[4], first[5], first[6]),
                         (1, "蝦皮門市", "板橋長江 - 智取店", "王小明", "A123456789", "115.09.21當天加退"))
        self.assertEqual(rows[-1][10], "離店")


class ShopeeUploadRouteTests(unittest.TestCase):
    def setUp(self):
        self.db = FakeFirestore()
        self.blobs = {}
        self.account = HR
        for p in (
            mock.patch.object(repo, "insurance_uploads_ref", side_effect=lambda: self.db.collection("u")),
            mock.patch.object(repo, "insurance_day_locks_ref", side_effect=lambda: self.db.collection("l")),
            mock.patch("hr.insurance_draft_repository.insurance_drafts_ref", side_effect=lambda: self.db.collection("d")),
            mock.patch.object(insurance_routes, "upload_file", side_effect=self._upload),
            mock.patch.object(platform_accounts, "current_account", side_effect=lambda request: self.account),
        ):
            p.start()
            self.addCleanup(p.stop)
        self.client = TestClient(main.app)

    def _upload(self, category, entity_id, filename, content, content_type):
        path = f"hr/{category}/{entity_id}/{len(self.blobs)}.xlsx"
        self.blobs[path] = content
        return path

    def _post(self, kind, content):
        return self.client.post(
            "/hr/insurance/upload",
            data={"work_date": "2026-09-25", "department": "蝦皮", "kind": kind},
            files={"file": ("f.xlsx", content, XLSX)},
            follow_redirects=False,
        )

    def test_two_kinds_are_stored_separately(self):
        self.assertEqual(self._post("elearning", elearning_file()).status_code, 303)
        self.assertEqual(self._post("notice", notice_file()).status_code, 303)
        e = repo.get_upload("蝦皮", "2026-09-25", "elearning")
        n = repo.get_upload("蝦皮", "2026-09-25", "notice")
        self.assertEqual((e["kind"], n["kind"]), ("elearning", "notice"))
        self.assertIn("E-learning 共 2 筆", e["summary"])
        html = self.client.get("/hr/insurance/summary?work_date=2026-09-25").text
        self.assertIn("蝦皮（E-learning）", html)
        self.assertIn("蝦皮（離店與實習通報）", html)
        self.assertEqual(html.count("已上傳"), 2)

    def test_wrong_file_is_rejected_with_message(self):
        response = self._post("notice", elearning_file())
        self.assertEqual(response.status_code, 400)
        self.assertIn("請確認上傳的檔案種類是否選對", response.text)
        self.assertIsNone(repo.get_upload("蝦皮", "2026-09-25", "notice"))

    def test_upload_page_shows_kind_picker(self):
        html = self.client.get("/hr/insurance/upload?department=蝦皮&kind=notice&work_date=2026-09-25").text
        self.assertIn('<option value="notice" selected>', html)
        self.assertIn("離店與實習通報 Excel 檔案", html)
        self.assertIn('<option value="蝦皮" selected>', html)

    def test_department_staff_cannot_upload_shopee(self):
        self.account = TAOYUAN
        self._post("elearning", elearning_file())
        self.assertIsNone(repo.get_upload("蝦皮", "2026-09-25", "elearning"))


if __name__ == "__main__":
    unittest.main()
