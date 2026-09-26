"""薪資補款核准信＋PDF 存查單由平台產生（2026-09-26）。測試資料全部是假的。"""
import io
import os
import sys
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests import _env  # noqa: F401
from tests import _stub_gcp
_stub_gcp.install()

from fastapi.testclient import TestClient

import finance_routes
import main
import platform_accounts
from services import salary_repayment_report as report
from services import salary_repayment_store as store
from tests._fake_firestore import FakeFirestore

# GAS 預設的 22 欄標題列
HEADERS = ["補款單號", "申請時間", "申請人姓名", "申請人 LINE ID", "員工姓名", "身分證", "廠商/店家", "申請日", "付款日",
           "扣分鐘月份", "補請款月份", "是否可請款", "補款方式", "加項小計", "扣項小計", "實補總額", "備註", "審核狀態",
           "核准主管", "核准時間", "補款佐證(照片)", "匯費"]
ORG_HEADERS = ["員工姓名", "員工 LINE ID", "主管姓名", "主管 LINE ID", "主管 Email", "F", "G", "H", "I", "員工Email"]
APPLICANT_ID = "Uapplicant0001"
MANAGER_ID = "Umanager00001"
ADMIN = {"username": "boss", "name": "老闆", "department": "", "is_platform_admin": True, "modules": []}


def fields(**overrides):
    values = ["SAL-20260901100000", "2026-09-01 10:00:00", "王小明", APPLICANT_ID, "陳小華", "A123456789", "測試店家",
              "2026-09-01", "", "", "2026-08", "是", "匯款", "1200", "30", "1170", "漏發加班費", "已核准",
              MANAGER_ID, "2026-09-02 09:00:00", "", "30"]
    data = dict(zip(HEADERS, values))
    data.update(overrides)
    return data


def org_rows():
    return [
        dict(zip(ORG_HEADERS, ["王小明", APPLICANT_ID, "李主管", MANAGER_ID, "li@example.com", "", "", "", "", "wang@example.com"])),
        dict(zip(ORG_HEADERS, ["李主管", MANAGER_ID, "", "", "", "", "", "", "", ""])),
    ]


class FormattingTests(unittest.TestCase):
    def test_minguo(self):
        self.assertEqual(report.format_minguo("2026-09-14", True), "115.09.14")
        self.assertEqual(report.format_minguo("2026/9/4", True), "115.09.04")
        self.assertEqual(report.format_minguo("2026-08", False), "115.08")
        self.assertEqual(report.format_minguo("看不懂的日期", True), "看不懂的日期")
        self.assertEqual(report.format_minguo("", True), "")

    def test_money_like_to_locale_string(self):
        self.assertEqual(report.format_money("1200"), "1,200")
        self.assertEqual(report.format_money("1,200"), "1,200")
        self.assertEqual(report.format_money(""), "0")
        self.assertEqual(report.format_money("12.5"), "12.5")

    def test_split_multi_like_gas(self):
        self.assertEqual([s for s in report.split_multi("a@x.com，b@x.com、c@x.com") if s], ["a@x.com", "b@x.com", "c@x.com"])


class RecordTests(unittest.TestCase):
    def setUp(self):
        self.org = report.Org(org_rows(), ORG_HEADERS)

    def test_fields_are_read_by_position_not_header_text(self):
        renamed = [h if i != 13 else "加項總額" for i, h in enumerate(HEADERS)]
        data = dict(zip(renamed, [fields()[h] for h in HEADERS]))
        record = report.build_record(data, renamed, self.org)
        self.assertEqual(record["total_earnings"], "1200")

    def test_emails_and_approver_name(self):
        record = report.build_record(fields(), HEADERS, self.org)
        self.assertEqual(record["supervisor_email"], "li@example.com")
        self.assertEqual(record["applicant_email"], "wang@example.com")
        self.assertEqual(record["approved_supervisor_name"], "李主管")
        self.assertEqual(report.recipients(record, "fin@example.com, li@example.com"),
                         ["fin@example.com", "li@example.com", "wang@example.com"])

    def test_unknown_approver_falls_back_to_raw_value(self):
        record = report.build_record(fields(**{"核准主管": "Uunknown00001"}), HEADERS, self.org)
        self.assertEqual(record["approved_supervisor_name"], "Uunknown00001")

    def test_supervisor_found_by_name_when_line_ids_missing(self):
        rows = org_rows()
        rows[0]["主管 LINE ID"] = ""
        sups = report.Org(rows, ORG_HEADERS).supervisors(APPLICANT_ID, "王小明")
        self.assertEqual(sups, [{"name": "李主管", "line_id": MANAGER_ID, "email": "li@example.com"}])

    def test_applicant_email_falls_back_to_their_supervisor_column(self):
        rows = org_rows()
        rows[0]["員工Email"] = ""
        rows.append(dict(zip(ORG_HEADERS, ["張三", "Uzhang0000001", "王小明", APPLICANT_ID, "wang-boss@example.com", "", "", "", "", ""])))
        self.assertEqual(report.Org(rows, ORG_HEADERS).applicant_email("王小明", APPLICANT_ID), "wang-boss@example.com")

    def test_subject_and_html(self):
        record = report.build_record(fields(**{"備註": "<script>x</script>"}), HEADERS, self.org)
        self.assertEqual(report.subject(record), "【薪資補款單 - 審核通過】陳小華 - 測試店家 (單號: SAL-20260901100000)")
        html = report.email_html(record, image_src="cid:salaryProofImg")
        for text in ("115.09.01", "115.08", "NT$ 1,200", "NT$ 1,170", "核准主管：李主管", "尚未指定", 'src="cid:salaryProofImg"'):
            self.assertIn(text, html)
        self.assertNotIn("<script>x</script>", html)
        self.assertNotIn("開啟 Google 雲端", html)  # 沒給原圖連結就不放按鈕
        self.assertNotIn("補款佐證單據與圖檔", report.email_html(record))

    def test_docx_contains_all_sections(self):
        import docx

        record = report.build_record(fields(), HEADERS, self.org)
        document = docx.Document(io.BytesIO(report.build_docx(record)))
        text = "\n".join(p.text for p in document.paragraphs)
        cells = [c.text for t in document.tables for r in t.rows for c in r.cells]
        for heading in ("薪資補款申請存查單", "一、基本資料與請款明細", "二、金額明細", "三、簽核紀錄"):
            self.assertIn(heading, text)
        for value in ("測試店家", "陳小華", "115.09.01", "NT$ 30", "漏發加班費", "李主管", "2026-09-02 09:00:00"):
            self.assertIn(value, cells)
        self.assertIn("NT$ 1,170", "\n".join(cells))


class PreviewRouteTests(unittest.TestCase):
    def setUp(self):
        self.db = FakeFirestore()
        self.account = ADMIN
        for p in (
            mock.patch.object(store, "get_db", return_value=self.db),
            mock.patch.object(platform_accounts, "current_account", side_effect=lambda request: self.account),
            mock.patch.object(report, "SALARY_HR_ACCOUNTING_EMAILS", "fin@example.com"),
        ):
            p.start()
            self.addCleanup(p.stop)
        org_values = [ORG_HEADERS] + [[r[h] for h in ORG_HEADERS] for r in org_rows()]
        pending = fields(**{"補款單號": "SAL-2", "審核狀態": "待審核"})
        store.sync_from_sheet(org_values, [HEADERS, [fields()[h] for h in HEADERS], [pending[h] for h in HEADERS]], ADMIN)
        self.client = TestClient(main.app)

    def test_list_shows_only_approved(self):
        html = self.client.get("/finance/migration").text
        self.assertIn("/finance/migration/preview/SAL-20260901100000", html)
        self.assertNotIn("/finance/migration/preview/SAL-2\"", html)

    def test_preview_page(self):
        resp = self.client.get("/finance/migration/preview/SAL-20260901100000")
        self.assertEqual(resp.status_code, 200)
        self.assertIn("fin@example.com、li@example.com、wang@example.com", resp.text)
        self.assertIn("只是預覽，沒有寄出", resp.text)
        self.assertIn("核准主管：李主管", resp.text)

    def test_pdf_download(self):
        with mock.patch.object(report, "build_pdf", return_value=b"%PDF-fake") as build:
            resp = self.client.get("/finance/migration/preview/SAL-20260901100000/pdf")
        self.assertEqual((resp.status_code, resp.content), (200, b"%PDF-fake"))
        self.assertEqual(build.call_args.args[0]["net_total"], "1170")

    def test_missing_and_admin_only(self):
        self.assertEqual(self.client.get("/finance/migration/preview/NOPE").status_code, 404)
        self.account = {**ADMIN, "is_platform_admin": False, "department": "財務部"}
        resp = self.client.get("/finance/migration/preview/SAL-20260901100000", follow_redirects=False)
        self.assertEqual(resp.status_code, 303)


if __name__ == "__main__":
    unittest.main()
