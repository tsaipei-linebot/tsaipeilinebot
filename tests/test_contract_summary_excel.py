import io
import os
import sys
import unittest

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests import _env  # noqa: F401
from tests import _stub_gcp
_stub_gcp.install()

from openpyxl import load_workbook

from services.contract_summary_excel import (
    build_client_contract_summary_workbook,
    build_dispatch_contract_summary_workbook,
    build_merged_summary_workbook,
)


class BuildClientContractSummaryWorkbookTests(unittest.TestCase):
    def test_header_and_row_content(self):
        rows = [{
            "client_name": "測試客戶", "tax_id": "12345678", "year": 2026, "party_b_name": "材霈甲公司",
            "contract_version": "hourly_flat_rate", "pricing_summary": "時薪 196／管理費 20",
            "remit_day": "10", "submitted_by": "alice",
        }]
        contract_versions = {"hourly_flat_rate": {"label": "時薪一口價"}}
        content = build_client_contract_summary_workbook(rows, contract_versions)

        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        row = [cell.value for cell in ws[2]]
        self.assertEqual(header, ["客戶名稱", "統一編號", "合約年", "簽約公司", "合約版本", "報價方式", "匯款截止日", "送出人"])
        self.assertEqual(row, ["測試客戶", "12345678", 2026, "材霈甲公司", "時薪一口價", "時薪 196／管理費 20", "10", "alice"])

    def test_unknown_version_falls_back_to_raw_code(self):
        rows = [{
            "client_name": "A", "tax_id": "", "year": 2026, "party_b_name": "",
            "contract_version": "some_unknown_code", "pricing_summary": "", "remit_day": "", "submitted_by": "",
        }]
        content = build_client_contract_summary_workbook(rows, {})
        wb = load_workbook(io.BytesIO(content))
        row = [cell.value for cell in wb.active[2]]
        self.assertEqual(row[4], "some_unknown_code")


class BuildDispatchContractSummaryWorkbookTests(unittest.TestCase):
    def test_header_and_row_content(self):
        rows = [{
            "client_name": "pchome", "submitted_by": "bob", "updated_year": 2026, "title": "日班", "hours": "9-18",
            "wage": "196", "bonus": "－", "overtime": "－", "pay_cycle": "每月10號",
        }]
        content = build_dispatch_contract_summary_workbook(rows)
        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        row = [cell.value for cell in ws[2]]
        self.assertEqual(header, ["客戶名稱", "送出人", "最近異動年份", "職稱/班別", "工作時間", "時薪", "工時獎金", "加班", "結薪週期"])
        self.assertEqual(row, ["pchome", "bob", 2026, "日班", "9-18", "196", "－", "－", "每月10號"])

    def test_none_updated_year_becomes_blank_cell(self):
        # openpyxl 存讀空字串會正規化成 None（實測過），所以這裡驗證的是
        # 「不是 0 或字面上的 'None' 這種難看的值」，不是驗證真的存成 ""。
        rows = [{
            "client_name": "pchome", "submitted_by": "bob", "updated_year": None, "title": "日班", "hours": "",
            "wage": "", "bonus": "", "overtime": "", "pay_cycle": "",
        }]
        content = build_dispatch_contract_summary_workbook(rows)
        wb = load_workbook(io.BytesIO(content))
        row = [cell.value for cell in wb.active[2]]
        self.assertIsNone(row[2])


class BuildMergedSummaryWorkbookTests(unittest.TestCase):
    def test_header_has_a_column_group_per_shift_slot(self):
        rows = [{
            "client_name": "A公司", "tax_id": "12345678", "year": 2026, "party_b_name": "材霈甲公司",
            "contract_version": "hourly_flat_rate", "pricing_summary": "時薪 196／管理費 20",
            "remit_day": "10", "submitted_by": "alice",
            "dispatch_pay_cycle": "每月10號", "dispatch_submitted_by": "bob",
            "shift_columns": [
                {"title": "日班", "hours": "9-18", "wage": "196", "bonus": "－", "overtime": "－"},
                {"title": "夜班", "hours": "22-7", "wage": "210", "bonus": "－", "overtime": "－"},
            ],
        }]
        content = build_merged_summary_workbook(rows, 2, {"hourly_flat_rate": {"label": "時薪一口價"}})
        wb = load_workbook(io.BytesIO(content))
        ws = wb.active
        header = [cell.value for cell in ws[1]]
        row = [cell.value for cell in ws[2]]
        self.assertEqual(header, [
            "客戶名稱", "統一編號", "合約年", "簽約公司", "合約版本", "報價方式", "匯款截止日", "合約送出人",
            "契約結薪週期", "契約送出人",
            "班別1-職稱", "班別1-工作時間", "班別1-時薪", "班別1-工時獎金", "班別1-加班",
            "班別2-職稱", "班別2-工作時間", "班別2-時薪", "班別2-工時獎金", "班別2-加班",
        ])
        self.assertEqual(row[:10], ["A公司", "12345678", 2026, "材霈甲公司", "時薪一口價", "時薪 196／管理費 20", "10", "alice", "每月10號", "bob"])
        self.assertEqual(row[10:15], ["日班", "9-18", "196", "－", "－"])
        self.assertEqual(row[15:20], ["夜班", "22-7", "210", "－", "－"])

    def test_no_shift_slots_still_produces_valid_header(self):
        content = build_merged_summary_workbook([], 0, {})
        wb = load_workbook(io.BytesIO(content))
        header = [cell.value for cell in wb.active[1]]
        self.assertEqual(len(header), 10)


if __name__ == "__main__":
    unittest.main()
