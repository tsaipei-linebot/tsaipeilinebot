import os
import sys
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests import _env  # noqa: F401
from tests import _stub_gcp
_stub_gcp.install()

import docx

from services.dispatch_contract_service import (
    CLAUSE_DEFAULTS,
    CLAUSE_ORDER,
    SHIFT_COLUMN_CODES,
    build_shift_rows,
    can_view_submission,
    convert_docx_to_pdf,
    list_visible_submissions,
    render_contract_docx,
)

_SYSTEM_TOKENS = [
    "${com_full_name}", "${rec_name}", "${rec_id}", "${cus_name}", "${rec_rdate}",
    "${pap_sign}", "${com_leader}", "${com_id}", "${com_address}",
]


class BuildShiftRowsTests(unittest.TestCase):
    """build_shift_rows() 是純函式（不碰 Firestore），涵蓋欄位勾選跟
    忽略空白列這兩條核心規則。"""

    def test_disabled_column_shows_dash_even_if_filled(self):
        rows = build_shift_rows(
            [{"title": "日班", "hours": "08:00-17:00", "wage": "依法定", "bonus": "", "overtime": ""}],
            enabled_columns=["title", "wage"],
        )
        self.assertEqual(rows, [{"title": "日班", "hours": "－", "wage": "依法定", "bonus": "－", "overtime": "－"}])

    def test_enabled_column_left_blank_shows_dash(self):
        rows = build_shift_rows(
            [{"title": "", "hours": "", "wage": "196/hr", "bonus": "", "overtime": ""}],
            enabled_columns=SHIFT_COLUMN_CODES,
        )
        self.assertEqual(rows[0]["title"], "－")
        self.assertEqual(rows[0]["wage"], "196/hr")

    def test_completely_empty_row_is_dropped(self):
        rows = build_shift_rows(
            [{"title": "", "hours": "", "wage": "", "bonus": "", "overtime": ""},
             {"title": "PT", "hours": "", "wage": "", "bonus": "", "overtime": ""}],
            enabled_columns=SHIFT_COLUMN_CODES,
        )
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["title"], "PT")

    def test_unknown_column_code_is_ignored(self):
        rows = build_shift_rows(
            [{"title": "日班", "hours": "", "wage": "", "bonus": "", "overtime": ""}],
            enabled_columns=["title", "not_a_real_column"],
        )
        self.assertEqual(rows, [{"title": "日班", "hours": "－", "wage": "－", "bonus": "－", "overtime": "－"}])


class RenderContractDocxTests(unittest.TestCase):
    """render_contract_docx() 會真的讀 assets/ 底下的 master template 檔案、
    用 docxtpl 套版，但不碰 Firestore／GCS，照既有分工歸類為可以直接測的
    純邏輯（檔案讀取是本地固定資源，不是外部服務）。"""

    def _render(self, **overrides):
        defaults = dict(
            work_address="台北市內湖區瑞光路588號",
            work_content="收銀結帳",
            pay_cycle="每月1號至月底",
            shifts=[{"title": "日班", "hours": "08:00-17:00", "wage": "依法定", "bonus": "每小時+津貼", "overtime": "依法規"}],
            clauses={},
        )
        defaults.update(overrides)
        content = render_contract_docx(**defaults)
        path = "/tmp/_test_dispatch_contract_render.docx"
        with open(path, "wb") as f:
            f.write(content)
        return docx.Document(path)

    def test_system_placeholder_tokens_are_untouched(self):
        doc = self._render()
        full_text = "\n".join(p.text for p in doc.paragraphs)
        for token in _SYSTEM_TOKENS:
            self.assertIn(token, full_text, f"{token} 不見了，不應該被套版動到")

    def test_variable_fields_are_substituted(self):
        doc = self._render(work_address="測試地址ABC", work_content="測試內容XYZ")
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("測試地址ABC", full_text)
        self.assertIn("測試內容XYZ", full_text)
        self.assertNotIn("{{", full_text)
        self.assertNotIn("{%", full_text)

    def test_shift_rows_appear_in_table(self):
        doc = self._render(shifts=[
            {"title": "日班", "hours": "08:00-17:00", "wage": "依法定", "bonus": "每小時+津貼", "overtime": "依法規"},
            {"title": "夜班", "hours": "20:00-05:00", "wage": "依法定", "bonus": "每小時+津貼", "overtime": "依法規"},
        ])
        table = doc.tables[0]
        rows_text = [[c.text for c in row.cells] for row in table.rows]
        self.assertIn(["日班", "08:00-17:00", "依法定", "每小時+津貼", "依法規"], rows_text)
        self.assertIn(["夜班", "20:00-05:00", "依法定", "每小時+津貼", "依法規"], rows_text)

    def test_missing_clause_falls_back_to_default(self):
        doc = self._render(clauses={})
        full_text = "\n".join(p.text for p in doc.paragraphs)
        for key in CLAUSE_ORDER:
            # 預設文字可能被 \n 拆成好幾段，這裡只確認開頭那一小段有出現。
            snippet = CLAUSE_DEFAULTS[key].split("\n")[0][:10]
            self.assertIn(snippet, full_text)

    def test_clause_override_replaces_default(self):
        doc = self._render(clauses={"benefits_note": "這是測試用的自訂福利說明ABC"})
        full_text = "\n".join(p.text for p in doc.paragraphs)
        self.assertIn("這是測試用的自訂福利說明ABC", full_text)
        self.assertNotIn(CLAUSE_DEFAULTS["benefits_note"], full_text)

    def test_pap_sign_boxes_keep_their_border(self):
        """${pap_sign} 出現的兩個地方（第二條後面、文件最後）在真實範本裡是
        用段落框線圍成一個方框，給另一套簽署系統貼簽名圖檔用——2026-09-12
        發現第一版整理 master template 時這個框線被拿掉了（文字沒變但視覺
        格式不見了），修正後這裡加一個回歸測試，確保之後改範本不會又不小心
        把框線拿掉。"""
        doc = self._render()
        full_xml = doc.element.xml
        self.assertEqual(full_xml.count("${pap_sign}"), 2)
        self.assertEqual(full_xml.count("<w:pBdr>"), 6)


class ClauseConfigConsistencyTests(unittest.TestCase):
    def test_clause_defaults_has_exactly_the_clause_order_keys(self):
        self.assertEqual(set(CLAUSE_DEFAULTS.keys()), set(CLAUSE_ORDER))


class ConvertDocxToPdfDelegatesToSharedModuleTests(unittest.TestCase):
    """convert_docx_to_pdf() 的實際轉檔邏輯搬到
    services/docx_pdf_conversion.py（跟 client_contract_service.py 共用，
    見 tests/test_docx_pdf_conversion.py 的完整涵蓋），這裡只驗證有正確
    委派過去、log 訊息前綴符合這個功能的用語。"""

    def test_delegates_to_shared_conversion_with_dispatch_contract_log_prefix(self):
        with mock.patch("services.dispatch_contract_service._convert_docx_to_pdf",
                         return_value=b"%PDF-FAKE") as mock_convert:
            result = convert_docx_to_pdf(b"fake docx bytes")
        self.assertEqual(result, b"%PDF-FAKE")
        mock_convert.assert_called_once_with(b"fake docx bytes", log_prefix="[派遣契約 PDF 轉檔失敗]")


class CanViewSubmissionTests(unittest.TestCase):
    """can_view_submission()：送出者本人／送出者的主管／平台管理員可以看，
    其他跟這筆紀錄無關的帳號看不到（2026-09-11 依使用者要求收斂權限，見
    services/dispatch_contract_service.py 開頭說明）。"""

    def test_submitter_can_view_own_record(self):
        viewer = {"username": "bob", "is_platform_admin": False}
        record = {"submitted_by": "bob"}
        self.assertTrue(can_view_submission(viewer, record))

    def test_platform_admin_can_view_anyone(self):
        viewer = {"username": "boss", "is_platform_admin": True}
        record = {"submitted_by": "alice"}
        self.assertTrue(can_view_submission(viewer, record))

    def test_manager_of_submitter_can_view(self):
        viewer = {"username": "carol", "is_platform_admin": False}
        record = {"submitted_by": "alice"}
        with mock.patch("services.dispatch_contract_service.platform_accounts.get_account",
                         return_value={"username": "alice", "manager_usernames": ["carol"]}):
            self.assertTrue(can_view_submission(viewer, record))

    def test_unrelated_account_cannot_view(self):
        viewer = {"username": "dave", "is_platform_admin": False}
        record = {"submitted_by": "alice"}
        with mock.patch("services.dispatch_contract_service.platform_accounts.get_account",
                         return_value={"username": "alice", "manager_usernames": ["carol"]}):
            self.assertFalse(can_view_submission(viewer, record))

    def test_missing_submitter_account_cannot_view(self):
        viewer = {"username": "dave", "is_platform_admin": False}
        record = {"submitted_by": "someone_deleted"}
        with mock.patch("services.dispatch_contract_service.platform_accounts.get_account", return_value=None):
            self.assertFalse(can_view_submission(viewer, record))


class ListVisibleSubmissionsTests(unittest.TestCase):
    """list_visible_submissions()：平台管理員看全部；其他帳號只看得到自己
    送出的、或自己是送出者主管的那些紀錄。"""

    def test_platform_admin_sees_all(self):
        viewer = {"username": "boss", "is_platform_admin": True}
        records = [{"submitted_by": "alice"}, {"submitted_by": "bob"}]
        with mock.patch("services.dispatch_contract_service.list_submissions", return_value=records):
            self.assertEqual(list_visible_submissions(viewer), records)

    def test_non_admin_only_sees_own_and_subordinates(self):
        viewer = {"username": "carol", "is_platform_admin": False}
        records = [
            {"id": "1", "submitted_by": "carol"},
            {"id": "2", "submitted_by": "alice"},
            {"id": "3", "submitted_by": "dave"},
        ]

        def _fake_get_account(username):
            accounts = {
                "alice": {"username": "alice", "manager_usernames": ["carol"]},
                "dave": {"username": "dave", "manager_usernames": ["someone_else"]},
            }
            return accounts.get(username)

        with mock.patch("services.dispatch_contract_service.list_submissions", return_value=records):
            with mock.patch("services.dispatch_contract_service.platform_accounts.get_account",
                             side_effect=_fake_get_account):
                visible = list_visible_submissions(viewer)
        self.assertEqual([r["id"] for r in visible], ["1", "2"])


if __name__ == "__main__":
    unittest.main()
