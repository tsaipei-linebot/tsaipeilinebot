import datetime
import io
import os
import sys
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests import _env  # noqa: F401
from tests import _stub_gcp
_stub_gcp.install()

from openpyxl import Workbook

import hr.insurance_repository as repo
from hr.insurance_excel import (
    _insured_date_summary,
    _to_roc_date,
    build_summary_workbook,
    parse_department_workbook,
)


def _snapshot(doc_id, data, exists=True):
    snapshot = mock.Mock(exists=exists)
    snapshot.id = doc_id
    snapshot.to_dict.return_value = data
    return snapshot


class CanUploadTests(unittest.TestCase):
    def test_upload_department_can_upload(self):
        self.assertTrue(repo.can_upload({"department": "桃園所"}))

    def test_non_upload_department_cannot_upload(self):
        self.assertFalse(repo.can_upload({"department": "財務部"}))

    def test_collector_department_is_not_an_upload_department(self):
        self.assertFalse(repo.can_upload({"department": "人資部門"}))

    def test_fullwidth_parens_in_account_department_still_matches(self):
        """2026-09-22 新增：帳號的部門欄位如果是用全形括號打的「新北所
        （配送組）」，也要能對到 hr/config.py 半形括號的「新北所(配送組)」
        ——見 platform_accounts.normalize_department() 的說明。"""
        self.assertTrue(repo.can_upload({"department": "新北所（配送組）"}))


class IsCollectorTests(unittest.TestCase):
    def test_hr_department_is_collector(self):
        self.assertTrue(repo.is_collector({"department": "人資部門"}))

    def test_platform_admin_is_always_collector(self):
        self.assertTrue(repo.is_collector({"department": "桃園所", "is_platform_admin": True}))

    def test_upload_department_is_not_collector(self):
        self.assertFalse(repo.is_collector({"department": "桃園所"}))


class HasInsuranceAccessTests(unittest.TestCase):
    def test_unrelated_department_has_no_access(self):
        self.assertFalse(repo.has_insurance_access({"department": "財務部"}))

    def test_upload_department_has_access(self):
        self.assertTrue(repo.has_insurance_access({"department": "台中所"}))

    def test_collector_has_access(self):
        self.assertTrue(repo.has_insurance_access({"department": "人資部門"}))


class SaveAndGetUploadTests(unittest.TestCase):
    def test_save_upload_writes_expected_fields(self):
        fake_doc_ref = mock.Mock()
        fake_collection = mock.Mock()
        fake_collection.document.return_value = fake_doc_ref
        with mock.patch.object(repo, "insurance_uploads_ref", return_value=fake_collection):
            repo.save_upload("桃園所", "2026-09-22", "hr/insurance/x/y.xlsx", "桃園所0922.xlsx", "u1", "王小明")
        fake_collection.document.assert_called_once_with("2026-09-22__桃園所")
        payload = fake_doc_ref.set.call_args.args[0]
        self.assertEqual(payload["department"], "桃園所")
        self.assertEqual(payload["work_date"], "2026-09-22")
        self.assertEqual(payload["blob_path"], "hr/insurance/x/y.xlsx")
        self.assertEqual(payload["uploaded_by"], "u1")
        self.assertEqual(payload["uploaded_by_name"], "王小明")

    def test_get_upload_returns_none_when_missing(self):
        fake_doc_ref = mock.Mock()
        fake_doc_ref.get.return_value = _snapshot("id1", {}, exists=False)
        fake_collection = mock.Mock()
        fake_collection.document.return_value = fake_doc_ref
        with mock.patch.object(repo, "insurance_uploads_ref", return_value=fake_collection):
            self.assertIsNone(repo.get_upload("桃園所", "2026-09-22"))

    def test_get_upload_returns_data_with_id(self):
        fake_doc_ref = mock.Mock()
        fake_doc_ref.get.return_value = _snapshot("2026-09-22__桃園所", {"department": "桃園所"})
        fake_collection = mock.Mock()
        fake_collection.document.return_value = fake_doc_ref
        with mock.patch.object(repo, "insurance_uploads_ref", return_value=fake_collection):
            result = repo.get_upload("桃園所", "2026-09-22")
        self.assertEqual(result["id"], "2026-09-22__桃園所")
        self.assertEqual(result["department"], "桃園所")


class ListHistoryTests(unittest.TestCase):
    def test_list_department_history_sorts_newest_first(self):
        fake_collection = mock.Mock()
        fake_query = mock.Mock()
        fake_query.stream.return_value = [
            _snapshot("a", {"department": "桃園所", "work_date": "2026-09-14"}),
            _snapshot("b", {"department": "桃園所", "work_date": "2026-09-21"}),
        ]
        fake_collection.where.return_value = fake_query
        with mock.patch.object(repo, "insurance_uploads_ref", return_value=fake_collection):
            result = repo.list_department_history("桃園所")
        self.assertEqual([r["id"] for r in result], ["b", "a"])

    def test_list_all_history_filters_by_date_range(self):
        fake_collection = mock.Mock()
        fake_collection.stream.return_value = [
            _snapshot("a", {"department": "桃園所", "work_date": "2026-09-14"}),
            _snapshot("b", {"department": "台中所", "work_date": "2026-09-18"}),
            _snapshot("c", {"department": "高雄所", "work_date": "2026-09-21"}),
        ]
        with mock.patch.object(repo, "insurance_uploads_ref", return_value=fake_collection):
            result = repo.list_all_history(start_date="2026-09-15", end_date="2026-09-20")
        self.assertEqual([r["id"] for r in result], ["b"])


class DayLockTests(unittest.TestCase):
    def test_is_day_closed_false_when_no_record(self):
        fake_doc_ref = mock.Mock()
        fake_doc_ref.get.return_value = _snapshot("2026-09-22", {}, exists=False)
        fake_collection = mock.Mock()
        fake_collection.document.return_value = fake_doc_ref
        with mock.patch.object(repo, "insurance_day_locks_ref", return_value=fake_collection):
            self.assertFalse(repo.is_day_closed("2026-09-22"))

    def test_is_day_closed_true_when_closed(self):
        fake_doc_ref = mock.Mock()
        fake_doc_ref.get.return_value = _snapshot("2026-09-22", {"closed": True})
        fake_collection = mock.Mock()
        fake_collection.document.return_value = fake_doc_ref
        with mock.patch.object(repo, "insurance_day_locks_ref", return_value=fake_collection):
            self.assertTrue(repo.is_day_closed("2026-09-22"))

    def test_close_day_writes_expected_fields(self):
        fake_doc_ref = mock.Mock()
        fake_collection = mock.Mock()
        fake_collection.document.return_value = fake_doc_ref
        with mock.patch.object(repo, "insurance_day_locks_ref", return_value=fake_collection):
            repo.close_day("2026-09-22", "hr1", "人資小美")
        payload = fake_doc_ref.set.call_args.args[0]
        self.assertTrue(payload["closed"])
        self.assertEqual(payload["closed_by"], "hr1")
        self.assertEqual(payload["closed_by_name"], "人資小美")

    def test_can_upload_for_date_blocked_when_closed(self):
        with mock.patch.object(repo, "is_day_closed", return_value=True):
            self.assertFalse(repo.can_upload_for_date({"department": "桃園所"}, "2026-09-22"))

    def test_can_upload_for_date_allowed_when_open(self):
        with mock.patch.object(repo, "is_day_closed", return_value=False):
            self.assertTrue(repo.can_upload_for_date({"department": "桃園所"}, "2026-09-22"))

    def test_collector_can_upload_even_when_closed(self):
        with mock.patch.object(repo, "is_day_closed", return_value=True):
            self.assertTrue(repo.can_upload_for_date({"department": "人資部門"}, "2026-09-22"))

    def test_non_upload_department_cannot_upload_even_when_open(self):
        with mock.patch.object(repo, "is_day_closed", return_value=False):
            self.assertFalse(repo.can_upload_for_date({"department": "財務部"}, "2026-09-22"))


class SummaryForDateTests(unittest.TestCase):
    def test_lists_all_upload_departments_in_order(self):
        with mock.patch.object(repo, "get_upload", return_value=None):
            rows = repo.summary_for_date("2026-09-22")
        self.assertEqual([r["department"] for r in rows], [
            "台北所(派遣組)", "台北所(國際組)", "新北所(派遣組)", "新北所(配送組)",
            "桃園所", "台中所", "高雄所", "蝦皮",
        ])
        self.assertTrue(all(r["upload"] is None for r in rows[:7]))
        # 蝦皮（2026-09-25）：一列、兩種檔案各自一格
        self.assertEqual([k["kind_name"] for k in rows[7]["shopee"]], ["E-learning", "離店與實習通報"])


class RocDateTests(unittest.TestCase):
    def test_converts_gregorian_to_roc(self):
        self.assertEqual(_to_roc_date(datetime.date(2026, 9, 19)), "115.09.19")


class InsuredDateSummaryTests(unittest.TestCase):
    def test_same_day_insured_and_withdrawn(self):
        result = _insured_date_summary(datetime.date(2026, 9, 19), datetime.date(2026, 9, 19))
        self.assertEqual(result, "115.09.19當天加退")

    def test_different_days(self):
        result = _insured_date_summary(datetime.date(2026, 9, 1), datetime.date(2026, 9, 19))
        self.assertEqual(result, "115.09.01加保／115.09.19退保")

    def test_insured_only(self):
        result = _insured_date_summary(datetime.date(2026, 9, 19), None)
        self.assertEqual(result, "115.09.19加保")

    def test_withdrawn_only(self):
        result = _insured_date_summary(None, datetime.date(2026, 9, 19))
        self.assertEqual(result, "115.09.19退保")

    def test_both_blank(self):
        self.assertEqual(_insured_date_summary(None, None), "")

    def test_unparseable_text_passes_through(self):
        result = _insured_date_summary("當天加退", "")
        self.assertEqual(result, "當天加退")


def _build_department_workbook(rows) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.append(["編號", "廠商", "班別", "姓名", "身分證", "勞保加保日期", "勞保退保日期", "勞保追退日期", "健保加保月份", "眷屬健保", "備註"])
    for row in rows:
        ws.append(row)
    buffer = io.BytesIO()
    wb.save(buffer)
    return buffer.getvalue()


class ParseDepartmentWorkbookTests(unittest.TestCase):
    def test_parses_data_rows(self):
        content = _build_department_workbook([
            [1, "蝦皮", "早班", "陳大文", "A123456789", datetime.date(2026, 9, 19), datetime.date(2026, 9, 19), None, "", "", "備註1"],
        ])
        rows = parse_department_workbook(content)
        self.assertEqual(len(rows), 1)
        self.assertEqual(rows[0]["姓名"], "陳大文")
        self.assertEqual(rows[0]["廠商"], "蝦皮")

    def test_skips_blank_trailing_rows(self):
        content = _build_department_workbook([
            [1, "蝦皮", "早班", "陳大文", "A123456789", None, None, None, "", "", ""],
            [None, None, None, None, None, None, None, None, None, None, None],
        ])
        rows = parse_department_workbook(content)
        self.assertEqual(len(rows), 1)


class BuildSummaryWorkbookTests(unittest.TestCase):
    def test_flattens_uploads_with_blank_unavailable_columns(self):
        content = _build_department_workbook([
            [1, "蝦皮", "早班", "陳大文", "A123456789", datetime.date(2026, 9, 19), datetime.date(2026, 9, 19), None, "", "", "備註1"],
        ])
        uploads = [{"department": "桃園所", "blob_path": "hr/insurance/x/y.xlsx"}]
        with mock.patch("hr.insurance_excel.download_file", return_value=(content, "application/octet-stream")):
            workbook_bytes = build_summary_workbook(uploads)

        wb_out = __import__("openpyxl").load_workbook(io.BytesIO(workbook_bytes))
        ws = wb_out.active
        header = [c.value for c in ws[1]]
        self.assertEqual(
            header,
            ["編號", "投保單位", "廠商", "部門/店家", "姓名", "身分證字號", "投保日", "出生年月日", "勞退追退日期", "班次/級距", "備註", "招募人員"],
        )
        data_row = [c.value for c in ws[2]]
        # 空字串儲存格經 openpyxl 存檔再讀回會變成 None，這裡統一視為「空白」。
        self.assertEqual(data_row[0], 1)
        self.assertIn(data_row[1], (None, ""))  # 投保單位無資料來源
        self.assertEqual(data_row[2], "蝦皮")
        self.assertEqual(data_row[3], "桃園所")
        self.assertEqual(data_row[4], "陳大文")
        self.assertEqual(data_row[5], "A123456789")
        self.assertEqual(data_row[6], "115.09.19當天加退")
        self.assertIn(data_row[7], (None, ""))  # 出生年月日無資料來源
        self.assertIn(data_row[9], (None, ""))  # 班次/級距無資料來源
        self.assertEqual(data_row[10], "備註1")
        self.assertIn(data_row[11], (None, ""))  # 招募人員無資料來源

    def test_skips_upload_with_missing_blob(self):
        uploads = [{"department": "桃園所", "blob_path": ""}]
        workbook_bytes = build_summary_workbook(uploads)
        wb_out = __import__("openpyxl").load_workbook(io.BytesIO(workbook_bytes))
        ws = wb_out.active
        self.assertEqual(ws.max_row, 1)  # 只剩表頭

    def test_skips_upload_when_file_not_found_in_storage(self):
        uploads = [{"department": "桃園所", "blob_path": "hr/insurance/x/missing.xlsx"}]
        with mock.patch("hr.insurance_excel.download_file", return_value=(None, None)):
            workbook_bytes = build_summary_workbook(uploads)
        wb_out = __import__("openpyxl").load_workbook(io.BytesIO(workbook_bytes))
        self.assertEqual(wb_out.active.max_row, 1)


if __name__ == "__main__":
    unittest.main()
