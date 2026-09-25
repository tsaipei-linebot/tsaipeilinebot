"""每日加退保暫存區（2026-09-24 新增）：資料層、Excel、上傳頁的送出／下載／紀錄。

用記憶體版 Firestore 跑真的資料流程；GCS 上傳／下載用 dict 模擬。"""
import io
import os
import sys
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests import _env  # noqa: F401
from tests import _stub_gcp
_stub_gcp.install()

import openpyxl
from fastapi.testclient import TestClient

import main
import platform_accounts
from hr import insurance_draft_repository as drafts
from hr import insurance_repository as repo
from hr.insurance_excel import build_department_workbook, parse_department_workbook
from hr.routes import insurance_routes
from tests._fake_firestore import FakeFirestore

DEPT = "新北所(配送組)"
AMY = {"username": "amy", "name": "Amy", "department": DEPT, "modules": [], "is_platform_admin": False, "rank": ""}
BOB_OTHER_DEPT = {"username": "bob", "name": "Bob", "department": "桃園所", "modules": [], "is_platform_admin": False, "rank": ""}
HR = {"username": "hr", "name": "HR", "department": "人資部門", "modules": ["hr"], "is_platform_admin": False, "rank": ""}


class _Env:
    """共用：假 Firestore + 假 GCS。"""

    def start(self, test):
        self.db = FakeFirestore()
        self.blobs = {}
        patchers = [
            mock.patch.object(drafts, "insurance_drafts_ref", side_effect=lambda: self.db.collection("hr_insurance_drafts")),
            mock.patch.object(repo, "insurance_uploads_ref", side_effect=lambda: self.db.collection("hr_insurance_uploads")),
            mock.patch.object(repo, "insurance_day_locks_ref", side_effect=lambda: self.db.collection("hr_insurance_day_locks")),
            mock.patch.object(insurance_routes, "upload_file", side_effect=self._upload),
            mock.patch.object(insurance_routes, "download_file", side_effect=lambda path: (self.blobs.get(path), "x") if path in self.blobs else (None, None)),
        ]
        for p in patchers:
            p.start()
            test.addCleanup(p.stop)

    def _upload(self, category, entity_id, filename, content, content_type):
        path = f"hr/{category}/{entity_id}/{len(self.blobs)}.xlsx"
        self.blobs[path] = content
        return path


def _ids(department=DEPT):
    """每日加退保頁畫面上會勾選的那些（2026-09-25 起送出要帶勾選的 id）。"""
    return [d["id"] for d in drafts.list_pending(department)]


def _rows(content):
    return [(r["姓名"], r["勞保加保日期"], r["勞保退保日期"]) for r in parse_department_workbook(content)]


class DraftRepositoryTests(unittest.TestCase):
    def setUp(self):
        _Env().start(self)

    def test_add_edit_cancel_keeps_the_record_and_history(self):
        draft_id = drafts.add_draft("新北所（配送組）", {"name": "王小明", "insured_date": "2026-09-25"}, AMY, kind=drafts.KIND_ADD, personnel_id="p1")
        draft = drafts.get_draft(draft_id)
        self.assertEqual(draft["department"], DEPT)  # 全形括號也歸到同一個部門
        self.assertEqual(draft["status"], drafts.STATUS_PENDING)
        self.assertEqual(drafts.draft_type_name(draft), "加保")

        self.assertTrue(drafts.update_draft(draft_id, {"name": "王小明", "note": "晚班"}, AMY))
        self.assertTrue(drafts.cancel_draft(draft_id, AMY, "手動刪除"))
        self.assertFalse(drafts.update_draft(draft_id, {"note": "x"}, AMY))  # 取消後不能改

        draft = drafts.get_draft(draft_id)
        self.assertEqual(draft["status"], drafts.STATUS_CANCELLED)
        self.assertEqual([h["action"] for h in draft["history"]], ["created", "edited", "cancelled"])
        self.assertEqual(draft["history"][1]["note"], "備註")
        self.assertEqual(drafts.list_pending(DEPT), [])
        self.assertEqual(len(drafts.list_drafts(DEPT)), 1)  # 紀錄還在

    def test_cancel_pending_for_personnel_reports_already_sent(self):
        sent_id = drafts.add_draft(DEPT, {"name": "甲", "withdrawn_date": "2026-09-01"}, AMY, kind=drafts.KIND_REMOVE, personnel_id="p1")
        drafts.mark_sent([drafts.get_draft(sent_id)], "2026-09-01", AMY)
        result = drafts.cancel_pending_for_personnel("p1", drafts.KIND_REMOVE, AMY, "改回")
        self.assertEqual(result, {"cancelled": 0, "already_sent": True})

        drafts.add_draft(DEPT, {"name": "甲", "withdrawn_date": "2026-09-30"}, AMY, kind=drafts.KIND_REMOVE, personnel_id="p1")
        result = drafts.cancel_pending_for_personnel("p1", drafts.KIND_REMOVE, AMY, "改回")
        self.assertEqual(result, {"cancelled": 1, "already_sent": False})

    def test_type_names(self):
        self.assertEqual(drafts.draft_type_name({"insured_date": "a", "withdrawn_date": "b"}), "加退保")
        self.assertEqual(drafts.draft_type_name({"withdrawn_date": "b"}), "退保")
        self.assertEqual(drafts.draft_type_name({"recovery_date": "c"}), "追退")


class DepartmentWorkbookTests(unittest.TestCase):
    def test_round_trip_matches_the_department_template(self):
        content = build_department_workbook([
            {"姓名": "王小明", "廠商": "UD", "勞保加保日期": "2026-09-25"},
            {"姓名": "=cmd", "勞保退保日期": "2026-09-30"},
        ])
        ws = openpyxl.load_workbook(io.BytesIO(content)).active
        self.assertEqual([c.value for c in ws[1]][:6], ["編號", "廠商", "班別", "姓名", "身分證", "勞保加保日期"])
        self.assertEqual(ws["A3"].value, 2)
        self.assertEqual(ws["D3"].value, "'=cmd")  # 擋公式注入
        rows = parse_department_workbook(content)
        self.assertEqual(rows[0]["勞保加保日期"].date().isoformat(), "2026-09-25")  # 真的 Excel 日期


class DraftRouteTests(unittest.TestCase):
    def setUp(self):
        self.env = _Env()
        self.env.start(self)
        self.account = AMY
        patcher = mock.patch.object(platform_accounts, "current_account", side_effect=lambda request: self.account)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.client = TestClient(main.app)

    def _post(self, url, data):
        return self.client.post(url, data=data, follow_redirects=False)

    def _add(self, name, **dates):
        return self._post("/hr/insurance/drafts/new", {"work_date": "2026-09-25", "name": name, **dates})

    def test_upload_page_shows_pending_list(self):
        drafts.add_draft(DEPT, {"name": "王小明", "vendor": "UD", "insured_date": "2026-09-25"}, AMY, kind=drafts.KIND_ADD)
        html = self.client.get("/hr/insurance/upload?work_date=2026-09-25").text
        self.assertIn("2026-09-25 的名單（1 筆）", html)
        self.assertIn("配送系統（報到）", html)
        self.assertIn("送出給人資（2026-09-25）", html)

    def test_other_departments_do_not_get_the_staging_area(self):
        self.account = BOB_OTHER_DEPT
        html = self.client.get("/hr/insurance/upload").text
        self.assertNotIn("待送出清單", html)
        self.assertEqual(self._add("x", insured_date="2026-09-25").status_code, 303)
        self.assertEqual(self.env.db.docs("hr_insurance_drafts"), {})

    def test_manual_add_requires_name_and_a_date(self):
        self.assertEqual(self._add("", insured_date="2026-09-25").status_code, 400)
        self.assertIn("至少要填一個", self._add("王小明").text)
        self.assertIn("msg=", self._add("王小明", withdrawn_date="2026-09-25").headers["location"])
        self.assertEqual(drafts.list_pending(DEPT)[0]["kind"], drafts.KIND_MANUAL)

    def test_send_twice_same_day_keeps_everything_and_includes_manual_excel(self):
        # 早上手動上傳過一份 Excel
        manual = build_department_workbook([{"姓名": "手動的人", "勞保加保日期": "2026-09-25"}])
        self.env.blobs["hr/insurance/manual.xlsx"] = manual
        repo.save_upload(DEPT, "2026-09-25", "hr/insurance/manual.xlsx", "m.xlsx", "amy", "Amy")

        self._add("甲", insured_date="2026-09-25")
        self._post("/hr/insurance/drafts/send", {"work_date": "2026-09-25", "draft_ids": _ids()})
        self._add("乙", withdrawn_date="2026-09-25")
        response = self._post("/hr/insurance/drafts/send", {"work_date": "2026-09-25", "draft_ids": _ids()})
        self.assertIn("msg=", response.headers["location"])

        upload = repo.get_upload(DEPT, "2026-09-25")
        self.assertTrue(upload["generated_from_drafts"])
        self.assertEqual([h["mode"] for h in upload["upload_history"]], ["send", "send"])
        self.assertEqual(len(upload["draft_ids"]), 2)
        self.assertEqual([r[0] for r in _rows(self.env.blobs[upload["blob_path"]])], ["手動的人", "甲", "乙"])
        statuses = {d["name"]: (d["status"], d["sent_work_date"]) for d in drafts.list_drafts(DEPT)}
        self.assertEqual(statuses, {"甲": ("sent", "2026-09-25"), "乙": ("sent", "2026-09-25")})

    def test_send_blocked_after_closing_and_download_marks_downloaded(self):
        self._add("甲", insured_date="2026-09-25")
        repo.close_day("2026-09-25", "hr", "HR")
        response = self._post("/hr/insurance/drafts/send", {"work_date": "2026-09-25", "draft_ids": _ids()})
        self.assertIn("err=", response.headers["location"])
        self.assertIsNone(repo.get_upload(DEPT, "2026-09-25"))
        self.assertIn("下載勾選的資料", self.client.get("/hr/insurance/upload?work_date=2026-09-25").text)

        response = self._post("/hr/insurance/drafts/download", {"work_date": "2026-09-25", "draft_ids": _ids()})
        self.assertEqual(response.status_code, 200)
        self.assertEqual(_rows(response.content)[0][0], "甲")
        self.assertEqual(drafts.list_pending(DEPT), [])
        self.assertEqual(drafts.list_drafts(DEPT)[0]["status"], drafts.STATUS_DOWNLOADED)

    def test_manual_upload_warns_when_drafts_were_already_sent(self):
        self._add("甲", insured_date="2026-09-25")
        self._post("/hr/insurance/drafts/send", {"work_date": "2026-09-25", "draft_ids": _ids()})
        html = self.client.get("/hr/insurance/upload?work_date=2026-09-25").text
        self.assertIn("已經從待送出清單送出 1 筆", html)

    def test_edit_and_cancel_only_while_pending(self):
        self._add("甲", insured_date="2026-09-25")
        draft_id = drafts.list_pending(DEPT)[0]["id"]
        self._post(f"/hr/insurance/drafts/{draft_id}/edit", {"work_date": "2026-09-25", "name": "甲", "insured_date": "2026-09-26"})
        self.assertEqual(drafts.get_draft(draft_id)["insured_date"], "2026-09-26")
        self._post(f"/hr/insurance/drafts/{draft_id}/cancel", {"work_date": "2026-09-25"})
        self.assertEqual(drafts.get_draft(draft_id)["status"], drafts.STATUS_CANCELLED)
        response = self.client.get(f"/hr/insurance/drafts/{draft_id}/edit", follow_redirects=False)
        self.assertIn("err=", response.headers["location"])

    def test_cannot_touch_another_departments_draft(self):
        other_id = drafts.add_draft("桃園所", {"name": "別人", "insured_date": "2026-09-25"}, BOB_OTHER_DEPT)
        self._post(f"/hr/insurance/drafts/{other_id}/cancel", {"work_date": ""})
        self.assertEqual(drafts.get_draft(other_id)["status"], drafts.STATUS_PENDING)

    def test_records_page_and_export(self):
        self._add("甲", insured_date="2026-09-25")
        self._add("乙", withdrawn_date="2026-09-25")
        drafts.cancel_draft(drafts.list_pending(DEPT)[1]["id"], AMY, "手動刪除")
        html = self.client.get("/hr/insurance/drafts/records").text
        self.assertIn("甲", html)
        self.assertIn("已取消", html)
        self.assertIn("Amy 取消（手動刪除）", html)
        html = self.client.get("/hr/insurance/drafts/records?status=cancelled").text
        self.assertNotIn(">甲<", html)
        response = self.client.get("/hr/insurance/drafts/records?export=1")
        ws = openpyxl.load_workbook(io.BytesIO(response.content)).active
        self.assertEqual(ws.max_row, 3)

    def test_hr_sees_all_departments_records(self):
        drafts.add_draft(DEPT, {"name": "甲", "insured_date": "2026-09-25"}, AMY)
        self.account = HR
        html = self.client.get("/hr/insurance/drafts/records").text
        self.assertIn("甲", html)
        self.assertIn("全部部門", html)


XLSX = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"


class ProxyUploadTests(unittest.TestCase):
    """人資代傳（2026-09-25）：選部門，那天已經有檔案預設接在後面。"""

    def setUp(self):
        self.env = _Env()
        self.env.start(self)
        self.account = HR
        patcher = mock.patch.object(platform_accounts, "current_account", side_effect=lambda request: self.account)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.client = TestClient(main.app)

    def _upload(self, department, names, mode=None):
        data = {"work_date": "2026-09-25", "department": department}
        if mode:
            data["mode"] = mode
        content = build_department_workbook([{"姓名": n, "勞保加保日期": "2026-09-25"} for n in names])
        return self.client.post("/hr/insurance/upload", data=data, files={"file": ("a.xlsx", content, XLSX)}, follow_redirects=False)

    def _names(self, department):
        upload = repo.get_upload(department, "2026-09-25")
        return [r[0] for r in _rows(self.env.blobs[upload["blob_path"]])]

    def test_hr_sees_department_dropdown(self):
        html = self.client.get("/hr/insurance/upload?department=台中所&work_date=2026-09-25").text
        self.assertIn("每日加退保代傳", html)
        self.assertIn('<option value="台中所" selected>', html)
        self.assertNotIn("待送出清單（", html)

    def test_append_is_default_when_file_exists_and_replace_is_optional(self):
        self._upload("台中所", ["甲"])
        self._upload("台中所", ["乙"])
        self.assertEqual(self._names("台中所"), ["甲", "乙"])
        self._upload("台中所", ["丙"], mode="replace")
        self.assertEqual(self._names("台中所"), ["丙"])
        modes = [h["mode"] for h in repo.get_upload("台中所", "2026-09-25")["upload_history"]]
        self.assertEqual(modes, ["replace", "append", "replace"])

    def test_hr_can_upload_after_closing(self):
        repo.close_day("2026-09-25", "hr", "HR")
        self._upload("高雄所", ["甲"])
        self.assertEqual(self._names("高雄所"), ["甲"])

    def test_department_staff_cannot_choose_another_department(self):
        self.account = AMY
        self._upload("台中所", ["甲"])
        self.assertIsNone(repo.get_upload("台中所", "2026-09-25"))
        self.assertEqual(self._names(DEPT), ["甲"])

    def test_append_keeps_rows_sent_from_drafts(self):
        self.account = AMY
        self.client.post("/hr/insurance/drafts/new", data={"work_date": "2026-09-25", "name": "配送甲", "insured_date": "2026-09-25"})
        self.client.post("/hr/insurance/drafts/send", data={"work_date": "2026-09-25", "draft_ids": _ids()})
        self.account = HR
        self._upload(DEPT, ["人資乙"])
        self.assertEqual(self._names(DEPT), ["配送甲", "人資乙"])
        self.assertEqual(len(repo.get_upload(DEPT, "2026-09-25")["draft_ids"]), 1)


class LateSubmissionTests(unittest.TestCase):
    """收單後補件：部門送出補件 → 人資收進或退件（2026-09-25）。"""

    def setUp(self):
        self.env = _Env()
        self.env.start(self)
        self.account = AMY
        patcher = mock.patch.object(platform_accounts, "current_account", side_effect=lambda request: self.account)
        patcher.start()
        self.addCleanup(patcher.stop)
        self.client = TestClient(main.app)
        repo.close_day("2026-09-25", "hr", "HR")

    def _post(self, url, data):
        return self.client.post(url, data=data, follow_redirects=False)

    def _add(self, name):
        self._post("/hr/insurance/drafts/new", {"work_date": "2026-09-25", "name": name, "withdrawn_date": "2026-09-25"})

    def _late_ids(self):
        return [d["id"] for d in drafts.list_late(DEPT)]

    def test_closed_day_offers_late_submission(self):
        self._add("甲")
        html = self.client.get("/hr/insurance/upload?work_date=2026-09-25").text
        self.assertIn("送出補件給人資（2026-09-25）", html)
        self.assertIn("備用：下載勾選的資料", html)
        self.assertIn("err=", self._post("/hr/insurance/drafts/send", {"work_date": "2026-09-25", "draft_ids": _ids()}).headers["location"])

    def test_late_submit_and_withdraw(self):
        self._add("甲")
        self._post("/hr/insurance/drafts/late-submit", {"work_date": "2026-09-25", "draft_ids": _ids()})
        self.assertEqual(drafts.list_pending(DEPT), [])
        [late_id] = self._late_ids()
        self.assertIn("補件待收（1 筆）", self.client.get("/hr/insurance/upload?work_date=2026-09-25").text)
        self._post(f"/hr/insurance/drafts/{late_id}/late-withdraw", {"work_date": "2026-09-25"})
        self.assertEqual(drafts.get_draft(late_id)["status"], drafts.STATUS_PENDING)

    def test_late_submit_not_allowed_before_closing(self):
        self._add("甲")
        response = self._post("/hr/insurance/drafts/late-submit", {"work_date": "2026-09-26", "draft_ids": _ids()})
        self.assertIn("err=", response.headers["location"])
        self.assertEqual(self._late_ids(), [])

    def test_hr_accept_appends_to_that_days_file(self):
        self._add("甲")
        self._add("乙")
        self._post("/hr/insurance/drafts/late-submit", {"work_date": "2026-09-25", "draft_ids": _ids()})
        self.account = HR
        html = self.client.get("/hr/insurance/summary?work_date=2026-09-25").text
        self.assertIn("有 2 筆補件待處理", html)
        ids = self._late_ids()
        response = self._post("/hr/insurance/late/accept", {"work_date": "2026-09-25", "draft_ids": ids[:1]})
        self.assertIn("msg=", response.headers["location"])
        upload = repo.get_upload(DEPT, "2026-09-25")
        self.assertEqual([r[0] for r in _rows(self.env.blobs[upload["blob_path"]])], ["甲"])
        accepted = drafts.get_draft(ids[0])
        self.assertEqual((accepted["status"], accepted["sent_work_date"]), ("sent", "2026-09-25"))
        self.assertEqual(accepted["history"][-1]["action"], "accepted")
        self.assertEqual(accepted["history"][-1]["by_name"], "HR")
        self.assertEqual(accepted["created_by_name"], "Amy")  # 建立人員不變

    def test_hr_reject_returns_to_pending_with_reason_until_next_send(self):
        self._add("甲")
        self._post("/hr/insurance/drafts/late-submit", {"work_date": "2026-09-25", "draft_ids": _ids()})
        self.account = HR
        self._post("/hr/insurance/late/reject", {"work_date": "2026-09-25", "draft_ids": self._late_ids(), "reason": ""})
        self.account = AMY
        [draft] = drafts.list_pending(DEPT)
        self.assertEqual(draft["rejected_reason"], "已超過下班時間，請明天再送")
        self.assertIn("人資退件：已超過下班時間，請明天再送", self.client.get("/hr/insurance/upload?work_date=2026-09-26").text)
        self._post("/hr/insurance/drafts/send", {"work_date": "2026-09-26", "draft_ids": _ids()})
        sent = drafts.get_draft(draft["id"])
        self.assertEqual((sent["status"], sent["rejected_reason"]), ("sent", ""))

    def test_department_staff_cannot_accept(self):
        self._add("甲")
        self._post("/hr/insurance/drafts/late-submit", {"work_date": "2026-09-25", "draft_ids": _ids()})
        self._post("/hr/insurance/late/accept", {"work_date": "2026-09-25", "draft_ids": self._late_ids()})
        self.assertEqual(len(self._late_ids()), 1)

    def test_personnel_revert_also_cancels_late_submissions(self):
        draft_id = drafts.add_draft(DEPT, {"name": "甲", "withdrawn_date": "2026-09-25"}, AMY, kind=drafts.KIND_REMOVE, personnel_id="p1")
        drafts.submit_late([drafts.get_draft(draft_id)], "2026-09-25", AMY)
        result = drafts.cancel_pending_for_personnel("p1", drafts.KIND_REMOVE, AMY, "改回")
        self.assertEqual(result["cancelled"], 1)
        self.assertEqual(drafts.get_draft(draft_id)["status"], drafts.STATUS_CANCELLED)


if __name__ == "__main__":
    unittest.main()
