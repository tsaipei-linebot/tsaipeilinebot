import io
import os
import sys
import unittest
from unittest import mock

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from tests import _env  # noqa: F401
from tests import _stub_gcp

_stub_gcp.install()

from salesdev import excel_export, pipeline, repository, sheet_import
from salesdev.scrapers.base import JobLead
from tests._fake_firestore import FakeFirestore


class _FakeDbMixin:
    def setUp(self):
        self.db = FakeFirestore()
        patcher = mock.patch.object(repository, "get_db", return_value=self.db)
        patcher.start()
        self.addCleanup(patcher.stop)


def _fake_scraper(leads=None, error=None):
    module = mock.Mock()
    if error:
        module.collect_leads.side_effect = error
    else:
        module.collect_leads.return_value = leads or []
    return module


class PipelineTests(_FakeDbMixin, unittest.TestCase):
    def _lead(self, job_id):
        return JobLead(
            source="104", job_id=job_id, job_title="包裝員", company_name="天泰人力銀行",
            job_url=f"https://www.104.com.tw/job/{job_id}", work_address="桃園市龜山區華亞二路",
        )

    def test_one_source_failing_does_not_stop_the_others(self):
        scrapers = (
            ("104", _fake_scraper([self._lead("a1"), self._lead("a2")])),
            ("1111", _fake_scraper(error=RuntimeError("HTTP 403"))),
            ("chickpt", _fake_scraper([])),
        )
        with mock.patch.object(pipeline, "SCRAPERS", scrapers):
            summary = pipeline.run_daily_scrape()
        self.assertEqual(summary["source_counts"], {"104": 2, "1111": 0, "chickpt": 0})
        self.assertEqual(summary["new_jobs"], 2)
        self.assertEqual(summary["new_groups"], 1)
        self.assertTrue(any("1111" in e and "403" in e for e in summary["errors"]))
        # 執行結果有記錄下來，畫面上方才能顯示「最近一次抓取」
        self.assertEqual(repository.latest_run()["new_jobs"], 2)

    def test_deadline_skips_remaining_sources(self):
        scrapers = (("104", _fake_scraper([self._lead("a1")])), ("chickpt", _fake_scraper([self._lead("b1")])))
        with mock.patch.object(pipeline, "SCRAPERS", scrapers), mock.patch.object(pipeline.time, "monotonic", side_effect=[0, 0, 999, 999, 999, 999]):
            summary = pipeline.run_daily_scrape(time_budget_seconds=10)
        self.assertTrue(summary["deadline_hit"])
        scrapers[1][1].collect_leads.assert_not_called()
        self.assertEqual(summary["new_jobs"], 1)

    def test_line_push_only_when_target_configured_and_new_jobs(self):
        bot = mock.Mock()
        scrapers = (("104", _fake_scraper([self._lead("a1")])),)
        with mock.patch.object(pipeline, "SCRAPERS", scrapers), mock.patch.object(pipeline, "SALESDEV_LINE_TARGET_ID", ""):
            pipeline.run_daily_scrape(bot)
        bot.push_message.assert_not_called()
        with mock.patch.object(pipeline, "SCRAPERS", (("104", _fake_scraper([self._lead("a9")])),)), mock.patch.object(pipeline, "SALESDEV_LINE_TARGET_ID", "Uxxx"):
            summary = pipeline.run_daily_scrape(bot)
        self.assertTrue(summary["line_pushed"])
        bot.push_message.assert_called_once()


class SheetRowTests(unittest.TestCase):
    def test_before_cutover_phone_goes_to_legacy_and_own_company_is_stripped(self):
        lead, legacy = sheet_import.lead_from_sheet_row(
            {
                "抓取日期": "2026-09-06", "來源平台": "chickpt", "公司名稱": "悅盛人力資源有限公司",
                "職缺名稱": "AI大廠作業員", "電話": "01-053-190", "工作地址": "台灣桃園市蘆竹區南青路XXXX號",
                "推測要派公司": "材霈有限公司、悅盛人力資源有限公司", "職缺連結": "https://www.chickpt.com.tw/job-OB31x3L2ym9l",
            }
        )
        self.assertEqual(lead["phone"], "")
        self.assertEqual(lead["seen_date"], "2026-09-06")
        self.assertEqual(legacy["phone"], "01-053-190")
        self.assertEqual(legacy["candidate_client_company"], "悅盛人力資源有限公司")

    def test_after_cutover_phone_is_the_posters(self):
        lead, legacy = sheet_import.lead_from_sheet_row(
            {"抓取日期": "2026-09-18", "來源平台": "104", "公司名稱": "天泰人力銀行", "電話": "989804029",
             "職缺連結": "https://www.104.com.tw/job/95wuw", "審查狀態": "待審查"}
        )
        self.assertEqual(lead["job_id"], "95wuw")
        self.assertEqual(lead["phone"], "989804029")
        self.assertNotIn("candidate_client_company", legacy)

    def test_rows_without_link_are_skipped(self):
        self.assertIsNone(sheet_import.lead_from_sheet_row({"來源平台": "104", "職缺連結": ""}))


class ImportFromSheetTests(_FakeDbMixin, unittest.TestCase):
    LEADS_HEADER = ["抓取日期", "來源平台", "公司名稱", "職缺名稱", "電話", "分機", "Email", "地區", "工作地址",
                    "推測要派公司", "要派公司比對備註", "職缺更新日期", "比對關鍵字", "職缺連結", "公司頁面連結", "審查狀態"]

    def _row(self, date, source, company, title, address, url, status=""):
        return [date, source, company, title, "", "", "", "", address, "", "", "", "", url, "", status]

    def test_import_groups_rows_and_is_repeatable(self):
        tabs = [
            ("新登記工廠", ["發現日期", "工廠名稱", "統一編號", "工廠地址", "行業別", "主要產品", "登記核准日期", "工廠登記編號"],
             [["2026-09-19", "新工廠", "12345678", "台中市西屯區", "", "", "1150901", ""]]),
            ("Leads", self.LEADS_HEADER, [
                self._row("2026-09-18", "chickpt", "悅盛人力資源有限公司", "日薪5940", "台灣桃園市桃園區桃鶯路xx號", "https://www.chickpt.com.tw/job-A1", "已勾選待反查"),
                self._row("2026-09-19", "chickpt", "悅盛人力資源有限公司", "黃仁勳沒發的", "台灣桃園市桃園區桃鶯路0號", "https://www.chickpt.com.tw/job-A2", "待審查"),
                self._row("2026-09-06", "104", "智邦人力資源管理顧問有限公司", "人力仲介行政助理", "", "https://www.104.com.tw/job/8b940"),
            ]),
        ]
        with mock.patch.object(sheet_import, "_read_all_tabs", return_value=tabs), mock.patch.object(sheet_import, "SALESDEV_SHEET_ID", "sheet"):
            stats = sheet_import.import_from_sheet("少凱")
            self.assertEqual(stats["error"], "")
            self.assertEqual((stats["jobs_read"], stats["new_jobs"], stats["new_groups"], stats["internal_jobs"]), (3, 3, 1, 1))
            self.assertEqual(stats["selected_groups"], 1)
            self.assertEqual(stats["new_factories"], 1)
            again = sheet_import.import_from_sheet("少凱")
        self.assertEqual((again["new_jobs"], again["updated_jobs"], again["new_factories"]), (0, 3, 0))
        [group] = repository.list_groups()
        self.assertEqual(group["review_status"], repository.STATUS_SELECTED)
        self.assertEqual(group["job_count"], 2)
        self.assertTrue(repository.get_job("chickpt_job-A1")["imported_from_sheet"])

    def test_read_failure_is_reported_in_plain_language(self):
        with mock.patch.object(sheet_import, "_read_all_tabs", side_effect=RuntimeError("403")), mock.patch.object(sheet_import, "SALESDEV_SHEET_ID", "sheet"):
            stats = sheet_import.import_from_sheet("少凱")
        self.assertIn("檢視者", stats["error"])


class ExcelExportTests(unittest.TestCase):
    def test_three_sheets_and_formula_injection_is_neutralised(self):
        import openpyxl

        groups = [{"id": "g1", "label": "桃園市桃園區桃鶯路", "review_status": "待審查", "job_count": 2, "agency_names": ["悅盛"],
                   "sources": ["chickpt"], "note": "=HYPERLINK(\"x\")", "contact_logs": [{"at": "2026-09-24 10:00", "by": "少凱", "text": "打過電話"}]}]
        jobs = [{"id": "j1", "group_id": "g1", "source": "chickpt", "job_title": "包裝員", "internal_reason": ""},
                {"id": "j2", "group_id": "", "source": "104", "job_title": "人力仲介行政", "internal_reason": "人力仲介"}]
        content = excel_export.build_workbook(groups, jobs, [{"name": "新工廠", "found_date": "2026-09-19"}])
        workbook = openpyxl.load_workbook(io.BytesIO(content))
        self.assertEqual(workbook.sheetnames, ["開發名單（依地點）", "全部職缺", "新登記工廠"])
        group_row = [c.value for c in workbook["開發名單（依地點）"][2]]
        self.assertEqual(group_row[14], "'=HYPERLINK(\"x\")")
        self.assertIn("打過電話", group_row[15])
        job_rows = {row[3].value: row for row in workbook["全部職缺"].iter_rows(min_row=2)}
        self.assertEqual(job_rows["包裝員"][0].value, "桃園市桃園區桃鶯路")
        self.assertEqual(job_rows["人力仲介行政"][10].value, "是")


if __name__ == "__main__":
    unittest.main()
